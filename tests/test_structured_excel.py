from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from app.services.structured_excel import (
    EPA_TEMPLATE_CODE,
    UNCERTAINTY_TEMPLATE_CODE,
    parse_epa_workbook,
    profile_workbook,
    sha256_file,
    write_parquet_tables,
)
from app.services.structured_query import StructuredFactorQuery, StructuredQueryService


def _append_epa_table(sheet, number: int) -> None:
    sheet.append([None, f"Table {number}", f"Table {number} title"])
    sheet.append([])
    if number == 1:
        sheet.append(
            [
                None,
                None,
                "Fuel Type",
                "Heat Content (HHV)",
                "CO2 Factor",
                "CH4 Factor",
                "N2O Factor",
                "CO2 Factor",
                "CH4 Factor",
                "N2O Factor",
            ]
        )
        sheet.append(
            [
                None,
                None,
                None,
                "mmBtu per short ton",
                "kg CO2 per mmBtu",
                "g CH4 per mmBtu",
                "g N2O per mmBtu",
                "kg CO2 per short ton",
                "g CH4 per short ton",
                "g N2O per short ton",
            ]
        )
        sheet.append([None, None, "Coal and Coke"])
        sheet.append([None, None, "Anthracite", 25.09, 103.69, 11, 1.6, 2602, 276, 40])
    elif number == 2:
        sheet.append([None, None, "Fuel Type", "kg CO2 per unit", "Unit"])
        sheet.append([None, None, "Motor Gasoline", 8.78, "gallon"])
    elif number == 3:
        sheet.append([None, None, "Vehicle Type", "Model Year", "CH4 Factor", "N2O Factor"])
        sheet.append([None, None, "Passenger Cars", "2020", 0.1, 0.2])
    elif number == 4:
        sheet.append(
            [None, None, "Vehicle Type", "Fuel Type", "Model Year", "CH4 Factor", "N2O Factor"]
        )
        sheet.append([None, None, "Passenger Cars", "Diesel", "2020", 0.1, 0.2])
    elif number == 5:
        sheet.append([None, None, "Vehicle Type", "Fuel Type", "CH4 Factor", "N2O Factor"])
        sheet.append([None, None, "Ships", "Diesel", 0.1, 0.2])
    elif number == 6:
        sheet.append([None, None, None, None, "Total Output", None, None, "Non-Baseload"])
        sheet.append(
            [
                None,
                None,
                "eGRID Subregion Acronym",
                "eGRID Subregion Name",
                "CO2 Factor",
                "CH4 Factor",
                "N2O Factor",
                "CO2 Factor",
                "CH4 Factor",
                "N2O Factor",
                "Grid Gross Loss (%)",
            ]
        )
        sheet.append(
            [None, None, None, None, "lb/MWh", "lb/MWh", "lb/MWh", "lb/MWh", "lb/MWh", "lb/MWh"]
        )
        sheet.append([None, None, "TEST", "Test Region", 1, 2, 3, 4, 5, 6, 0.04])
    elif number == 7:
        sheet.append([None, None, None, "CO2 Factor", "CH4 Factor", "N2O Factor"])
        sheet.append([None, None, "Steam and Heat", 1, 2, 3])
    elif number in {8, 10}:
        sheet.append(
            [None, None, "Vehicle Type", "CO2 Factor", "CH4 Factor", "N2O Factor", "Units"]
        )
        sheet.append([None, None, "Passenger Car", 1, 2, 3, "vehicle-mile"])
    elif number == 9:
        sheet.append(
            [
                None,
                None,
                "Material",
                "Recycled",
                "Landfilled",
                "Combusted",
                "Composted",
                "Anaerobically Digested Dry",
                "Anaerobically Digested Wet",
            ]
        )
        sheet.append([None, None, "Aluminum", 1, 2, 3, "NA", "NA", "NA"])
    elif number == 11:
        sheet.append(
            [
                None,
                None,
                "Industrial Designation or Common Name",
                "Chemical Formula",
                "100-Year GWP",
            ]
        )
        sheet.append([None, None, "Carbon dioxide", "CO2", 1])
    elif number == 12:
        sheet.append([None, None, "ASHRAE #", "100-year GWP", "Blend Composition"])
        sheet.append([None, None, "R-401A", 18, "blend"])
    sheet.append([None, None, "Source:", "Synthetic fixture"])


def _epa_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Emission Factors Hub"
    for number in range(1, 13):
        _append_epa_table(sheet, number)
    workbook.save(path)


def test_epa_template_extracts_twelve_typed_tables(tmp_path: Path) -> None:
    path = tmp_path / "2025 GHG Emission Factors Hub.xlsx"
    _epa_fixture(path)

    parsed = parse_epa_workbook(path, path.name)

    assert parsed.profile.template_code == EPA_TEMPLATE_CODE
    assert parsed.profile.edition_year == 2025
    assert len(parsed.tables) == 12
    assert {table.table_code for table in parsed.tables} == {
        "stationary_combustion",
        "mobile_combustion_co2",
        "mobile_onroad_gasoline",
        "mobile_onroad_alternative",
        "mobile_nonroad",
        "electricity",
        "steam_heat",
        "scope3_transportation",
        "scope3_waste",
        "scope3_travel_commuting",
        "gwp",
        "blended_refrigerant_gwp",
    }
    assert sum(len(table.records) for table in parsed.tables) == 38
    waste = next(table for table in parsed.tables if table.table_code == "scope3_waste")
    assert {record["applicability"] for record in waste.records} == {
        "AVAILABLE",
        "NOT_APPLICABLE",
    }
    assert all(record["source_row"] for table in parsed.tables for record in table.records)


def test_uncertainty_workbook_is_registered_but_not_treated_as_reference_data(
    tmp_path: Path,
) -> None:
    path = tmp_path / "uncertainty.xlsx"
    workbook = Workbook()
    workbook.active.title = "Instruction"
    workbook.create_sheet("Notes")
    workbook.create_sheet("Info")
    calculations = workbook.create_sheet("Calculations")
    calculations.cell(row=4, column=82, value="=1+1")
    workbook.save(path)

    profile = profile_workbook(path, path.name)

    assert profile.template_code == UNCERTAINTY_TEMPLATE_CODE
    assert profile.asset_type == "CALCULATION_MODEL"
    assert profile.formula_count == 1
    assert profile.details["activation_state"] == "MODEL_VALIDATION_REQUIRED"


def test_sha256_is_stable_for_duplicate_files(tmp_path: Path) -> None:
    first = tmp_path / "a.xlsx"
    second = tmp_path / "b.xlsx"
    _epa_fixture(first)
    second.write_bytes(first.read_bytes())

    assert sha256_file(first) == sha256_file(second)


def test_generated_parquet_supports_typed_parameterized_query(tmp_path: Path) -> None:
    path = tmp_path / "2025 GHG Emission Factors Hub.xlsx"
    _epa_fixture(path)
    parsed = parse_epa_workbook(path, path.name)
    files = write_parquet_tables(parsed, tmp_path / "parquet")

    rows = StructuredQueryService()._execute(
        [files["stationary_combustion"]],
        StructuredFactorQuery(
            dataset_ids=[1],
            table_code="stationary_combustion",
            activity="Anthracite",
            gas="CO2",
            denominator_unit="mmBtu",
        ),
    )

    assert len(rows) == 1
    assert rows[0]["factor_value"] == "103.690000000000"
    assert rows[0]["numerator_unit"] == "kg CO2"
    assert rows[0]["denominator_unit"] == "mmBtu"
    assert rows[0]["source_file"] == path.name

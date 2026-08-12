from __future__ import annotations

import base64
import zipfile
from io import BytesIO
from types import SimpleNamespace

import docx
import pytest
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml

from app.rag.core.markdown_parser.models import ElementType
from app.rag.core.markdown_parser.parser import MarkdownParser
from app.rag.core.parse_task_service import ParseTaskService
from app.rag.core.parser.exceptions import ParseBaseException
from app.rag.core.parser.providers.word_parser import WordParser
from app.services.document_ingestion import (
    DocumentQualityGateError,
    SimpleDocumentIngestionService,
)

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUB"
    "AScY42YAAAAASUVORK5CYII="
)


class _Storage:
    def __init__(self) -> None:
        self.uploads: list[dict] = []

    def upload_bytes(self, **kwargs) -> None:
        self.uploads.append(kwargs)

    @staticmethod
    def build_object_url(bucket: str, object_key: str) -> str:
        return f"http://minio.local/{bucket}/{object_key}"


def _complex_docx(tmp_path):
    document = docx.Document()
    document.add_heading("能碳核算表", level=1)
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "类别"
    table.cell(0, 1).text = "排放量"
    table.cell(0, 1).merge(table.cell(0, 2))
    table.cell(1, 0).text = "范围一"
    table.cell(1, 1).text = "2024"
    table.cell(1, 2).text = "2025"
    table.cell(2, 0).text = "燃料"
    table.cell(2, 1).text = "10"
    table.cell(2, 2).text = "9"
    nested = table.cell(2, 0).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "煤"
    nested.cell(0, 1).text = "5"
    document.add_picture(BytesIO(_PNG))
    path = tmp_path / "complex.docx"
    document.save(path)
    return path


def _quality_service() -> SimpleDocumentIngestionService:
    service = object.__new__(SimpleDocumentIngestionService)
    return service


def test_word_parser_persists_images_and_preserves_complex_table(tmp_path) -> None:
    source = _complex_docx(tmp_path)
    storage = _Storage()
    parser = WordParser(
        storage=storage,
        image_bucket="parsed",
        image_prefix="documents/1/images",
    )

    markdown = parser.parse(source)
    result = MarkdownParser().parse(markdown, source_file=source.name)

    assert len(storage.uploads) == 1
    assert "http://minio.local/parsed/documents/1/images/" in markdown
    assert "mock-minio://" not in markdown
    assert parser.metadata["image_assets_persisted"] is True
    assert parser.metadata["source_image_reference_count"] == 1
    assert parser.metadata["image_occurrence_count"] == 1
    assert parser.metadata["source_table_count"] == 2
    assert parser.metadata["source_top_level_table_count"] == 1
    assert parser.metadata["source_nested_table_count"] == 1
    assert parser.metadata["mammoth_html_table_count"] == 2
    assert len(result.tables) == 2
    assert "<table" not in markdown
    assert 'format="rag_text" schema="table-rag-v2"' in markdown
    nested_table = next(
        element
        for element in result.elements
        if element.metadata.get("parent_table_id") == "table-001"
    )
    assert nested_table.metadata["table_id"] == "table-001-001"
    previews = parser.metadata["word_table_previews"]
    assert [preview["id"] for preview in previews] == ["table-001", "table-001-001"]
    assert previews[0]["cells"][1]["column_span"] == 2
    assert "嵌套表格：table-001-001" in str(previews[0])
    assert "煤" in str(previews[1])


def test_word_quality_attaches_structure_and_passes_complete_docx(tmp_path) -> None:
    source = _complex_docx(tmp_path)
    parser = WordParser(
        storage=_Storage(),
        image_bucket="parsed",
        image_prefix="documents/1/images",
    )
    markdown = parser.parse(source)
    parse_result = MarkdownParser().parse(markdown, source_file=source.name)
    ParseTaskService._apply_word_page_markers(parse_result)
    parse_output = {
        "markdown": markdown,
        "parse_result": parse_result,
        "metadata": parser.extract_metadata(),
    }

    status, report = _quality_service()._process_word_quality(
        identity=SimpleNamespace(filename=source.name),
        parse_output=parse_output,
        markdown=markdown,
    )

    assert status == "PASSED"
    assert report["structured_table_count"] == 1
    assert report["structured_nested_table_count"] == 1
    table = next(e for e in parse_result.elements if e.type is ElementType.TABLE)
    assert table.metadata["table_structure"]["cells"][1]["column_span"] == 2
    assert table.metadata["page_number"] == 1


def test_word_quality_rejects_unpersisted_images(tmp_path) -> None:
    source = _complex_docx(tmp_path)
    parser = WordParser()
    markdown = parser.parse(source)
    parse_output = {
        "markdown": markdown,
        "parse_result": MarkdownParser().parse(markdown, source_file=source.name),
        "metadata": parser.extract_metadata(),
    }

    with pytest.raises(DocumentQualityGateError) as error:
        _quality_service()._process_word_quality(
            identity=SimpleNamespace(filename=source.name),
            parse_output=parse_output,
            markdown=markdown,
        )

    assert error.value.error_code == "WORD_CONTENT_VALIDATION_FAILED"
    assert "WORD_IMAGE_ASSETS_NOT_PERSISTED" in error.value.quality_report["blocking_issues"]


def test_word_quality_warns_but_does_not_block_unsupported_special_objects(
    tmp_path,
) -> None:
    document = docx.Document()
    document.add_paragraph("主体文本完整")
    source = tmp_path / "special-objects.docx"
    document.save(source)
    parser = WordParser()
    markdown = parser.parse(source)
    metadata = parser.extract_metadata()
    metadata.update(
        {
            "source_ole_object_count": 2,
            "source_ole_preview_count": 1,
            "source_chart_count": 1,
            "extracted_chart_count": 0,
            "special_object_warnings": ["An OLEObject was ignored"],
            "warnings": ["WORD_IMAGE_TRANSCODE_FAILED:image/x-wmf"],
        }
    )
    parse_output = {
        "markdown": markdown,
        "parse_result": MarkdownParser().parse(markdown, source_file=source.name),
        "metadata": metadata,
    }

    status, report = _quality_service()._process_word_quality(
        identity=SimpleNamespace(filename=source.name),
        parse_output=parse_output,
        markdown=markdown,
    )

    assert status == "PASSED"
    assert report["blocking_issues"] == []
    assert "WORD_IMAGE_TRANSCODE_FAILED:image/x-wmf" in report["warnings"]
    assert "WORD_OLE_PREVIEW_MISSING:source=2,preview=1" in report["warnings"]
    assert (
        "WORD_CHART_DATA_EXTRACTION_INCOMPLETE:source=1,extracted=0"
        in report["warnings"]
    )


def test_word_quality_warns_but_does_not_block_legacy_vml_images(tmp_path) -> None:
    source = _complex_docx(tmp_path)
    parser = WordParser(
        storage=_Storage(),
        image_bucket="parsed",
        image_prefix="documents/1/images",
    )
    markdown = parser.parse(source)
    metadata = parser.extract_metadata()
    metadata.update(
        {
            "source_image_reference_count": 4,
            "source_supported_image_reference_count": 1,
            "source_legacy_vml_image_reference_count": 3,
        }
    )
    parse_output = {
        "markdown": markdown,
        "parse_result": MarkdownParser().parse(markdown, source_file=source.name),
        "metadata": metadata,
    }

    status, report = _quality_service()._process_word_quality(
        identity=SimpleNamespace(filename=source.name),
        parse_output=parse_output,
        markdown=markdown,
    )

    assert status == "PASSED"
    assert report["blocking_issues"] == []
    assert report["source_image_reference_count"] == 4
    assert report["source_supported_image_reference_count"] == 1
    assert report["source_legacy_vml_image_reference_count"] == 3
    assert "WORD_LEGACY_VML_IMAGES_NOT_RENDERED:count=3" in report["warnings"]


def test_word_chart_cache_is_converted_to_retrieval_markdown() -> None:
    chart_xml = """<?xml version="1.0" encoding="UTF-8"?>
<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
 <c:chart><c:plotArea><c:lineChart><c:ser>
  <c:tx><c:strRef><c:strCache><c:pt idx="0"><c:v>排放量</c:v></c:pt></c:strCache></c:strRef></c:tx>
  <c:cat><c:strRef><c:strCache>
   <c:pt idx="0"><c:v>2024</c:v></c:pt><c:pt idx="1"><c:v>2025</c:v></c:pt>
  </c:strCache></c:strRef></c:cat>
  <c:val><c:numRef><c:numCache>
   <c:pt idx="0"><c:v>10</c:v></c:pt><c:pt idx="1"><c:v>8</c:v></c:pt>
  </c:numCache></c:numRef></c:val>
 </c:ser></c:lineChart></c:plotArea></c:chart>
</c:chartSpace>""".encode()
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("word/charts/chart1.xml", chart_xml)
    with zipfile.ZipFile(BytesIO(stream.getvalue())) as archive:
        charts = WordParser._extract_chart_summaries(archive, set(archive.namelist()))

    markdown = WordParser._render_chart_markdown(charts)

    assert charts[0]["type"] == "lineChart"
    assert "系列：排放量" in markdown
    assert "2024：10" in markdown
    assert "2025：8" in markdown


def test_word_parser_converts_formula_and_propagates_saved_pages(tmp_path) -> None:
    document = docx.Document()
    formula = document.add_paragraph()
    formula._p.append(
        parse_xml(
            '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            "<m:r><m:t>x</m:t></m:r><m:r><m:t>+</m:t></m:r>"
            "<m:r><m:t>1</m:t></m:r></m:oMath>"
        )
    )
    document.add_page_break()
    document.add_paragraph("第二页内容")
    source = tmp_path / "formula-pages.docx"
    document.save(source)

    parser = WordParser()
    markdown = parser.parse(source)
    parse_result = MarkdownParser().parse(markdown, source_file=source.name)
    marker_count = ParseTaskService._apply_word_page_markers(parse_result)

    assert "$x+1$" in markdown
    assert "<!-- WORD_PAGE:1 -->" in markdown
    assert "<!-- WORD_PAGE:2 -->" in markdown
    assert parser.metadata["source_formula_count"] == 1
    assert parser.metadata["converted_formula_count"] == 1
    assert parser.metadata["page_count"] == 2
    assert marker_count == 2
    second_page = next(e for e in parse_result.elements if "第二页内容" in e.content)
    assert second_page.metadata["page_number"] == 2


def test_word_parser_propagates_section_and_paragraph_page_breaks(tmp_path) -> None:
    document = docx.Document()
    document.add_paragraph("第一页内容")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("第二页内容")
    third_page = document.add_paragraph("第三页内容")
    third_page.paragraph_format.page_break_before = True
    source = tmp_path / "section-pages.docx"
    document.save(source)

    parser = WordParser()
    markdown = parser.parse(source)
    parse_result = MarkdownParser().parse(markdown, source_file=source.name)
    marker_count = ParseTaskService._apply_word_page_markers(parse_result)

    assert parser.metadata["page_count"] == 3
    assert parser.metadata["saved_page_break_count"] == 2
    assert parser.metadata["pagination_source"] == "saved_docx_pagination_markers"
    assert marker_count == 3
    first = next(e for e in parse_result.elements if "第一页内容" in e.content)
    second = next(e for e in parse_result.elements if "第二页内容" in e.content)
    third = next(e for e in parse_result.elements if "第三页内容" in e.content)
    assert first.metadata["page_number"] == 1
    assert second.metadata["page_number"] == 2
    assert third.metadata["page_number"] == 3


def test_word_parser_does_not_count_continuous_or_final_sections(tmp_path) -> None:
    document = docx.Document()
    document.add_paragraph("第一节")
    document.add_section(WD_SECTION.CONTINUOUS)
    document.add_paragraph("连续分节正文")
    source = tmp_path / "continuous-section.docx"
    document.save(source)

    parser = WordParser()
    markdown = parser.parse(source)

    assert parser.metadata["page_count"] == 1
    assert parser.metadata["saved_page_break_count"] == 0
    assert markdown.count("<!-- WORD_PAGE:") == 1


@pytest.mark.asyncio
async def test_parse_task_service_removes_word_page_markers_from_retrieval(tmp_path) -> None:
    document = docx.Document()
    document.add_paragraph("第一页")
    document.add_page_break()
    document.add_paragraph("第二页")
    source = tmp_path / "pages.docx"
    document.save(source)

    output = await ParseTaskService.aprocess(source, "docx", source_file=source.name)
    parse_result = output["parse_result"]

    assert output["metadata"]["word_page_markers"] == 2
    assert all("WORD_PAGE" not in element.content for element in parse_result.elements)
    first = next(element for element in parse_result.elements if "第一页" in element.content)
    second = next(element for element in parse_result.elements if "第二页" in element.content)
    assert first.metadata["page_number"] == 1
    assert second.metadata["page_number"] == 2


def test_word_parser_rejects_duplicate_ooxml_members(tmp_path) -> None:
    document = docx.Document()
    document.add_paragraph("正文")
    source = tmp_path / "duplicate-member.docx"
    document.save(source)
    with zipfile.ZipFile(source) as archive:
        document_xml = archive.read("word/document.xml")
    with pytest.warns(UserWarning, match="Duplicate name"):
        with zipfile.ZipFile(source, "a") as archive:
            archive.writestr("word/document.xml", document_xml)

    with pytest.raises(ParseBaseException, match="重复 ZIP 条目"):
        WordParser().parse(source)


def test_word_parser_rejects_invalid_raster_payload() -> None:
    with pytest.raises(ParseBaseException, match="无法安全解码"):
        WordParser._validate_raster_image(b"not-an-image")


def test_legacy_doc_conversion_requires_libreoffice(tmp_path, monkeypatch) -> None:
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy")
    monkeypatch.setattr(
        "app.rag.core.parser.providers.word_parser.shutil.which",
        lambda _binary: None,
    )

    with pytest.raises(ParseBaseException, match="未安装 LibreOffice"):
        WordParser()._convert_legacy_doc(source, tmp_path / "converted")

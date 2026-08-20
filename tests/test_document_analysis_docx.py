from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

import pytest
from docx import Document as WordDocument

from app.domain.models import Document
from app.rag.core.llm.response import UsageInfo
from app.services.document_analysis import AnalysisSource, DocumentAnalysisResult
from app.services.document_analysis_docx import build_document_analysis_docx


def test_document_analysis_docx_contains_cover_report_tables_and_source_index() -> None:
    source_document = Document(
        id=7,
        dataset_id=3,
        user_id=11,
        filename="企业碳足迹报告.docx",
        file_type="docx",
        file_size=1024,
        raw_bucket="raw",
        raw_object_key="raw/report.docx",
        parsed_bucket="private",
        parsed_object_key="parsed/report.md",
        status="READY",
        version=2,
    )
    result = DocumentAnalysisResult(
        markdown=(
            "# 企业文档分析报告\n\n"
            "## 1. 报告摘要\n\n"
            "**结论：**报告期已说明。[文档片段1]\n\n"
            "## 10. 文档规范性检查\n\n"
            "| 序号 | 检查项 | 结论 |\n"
            "| --- | --- | --- |\n"
            "| 1 | 报告期 | 满足 |\n\n"
            "## 12. 结论与分析限制\n\n"
            "- 不替代正式核查。"
        ),
        sources=[
            AnalysisSource(
                citation_index=1,
                chunk_id="chunk-1",
                chunk_index=0,
                chunk_type="mixed",
                page=3,
                page_range=None,
                excerpt="报告期为 2023 年。",
            )
        ],
        model_name="analysis-model",
        model_config_id=3,
        usage=UsageInfo(prompt_tokens=8, completion_tokens=4, total_tokens=12),
        analyzed_chunk_count=1,
        evidence_batch_count=1,
        generated_at=datetime(2026, 8, 20, tzinfo=UTC),
    )

    content = build_document_analysis_docx(document=source_document, result=result)
    assert content.startswith(b"PK")

    exported = WordDocument(BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in exported.paragraphs)
    assert "企业文档\n分析报告" in text
    assert "分析报告编号：AI000007-V2" in text
    assert "源文件：企业碳足迹报告.docx" in text
    assert "1. 报告摘要" in text
    assert "引用证据索引" in text
    assert len(exported.tables) == 2
    assert exported.tables[0].cell(1, 1).text == "报告期"
    assert exported.tables[1].cell(1, 0).text == "[文档片段1]"
    assert exported.tables[1].cell(1, 1).text == "第 3 页"
    assert exported.sections[0].page_width.cm == pytest.approx(21, abs=0.01)
    assert exported.sections[0].page_height.cm == pytest.approx(29.7, abs=0.01)
    assert exported.sections[0].left_margin.cm == pytest.approx(3.5, abs=0.01)
    assert exported.sections[0].right_margin.cm == pytest.approx(3.0, abs=0.01)
    assert exported.styles["Normal"].font.size.pt == 12
    assert exported.styles["Heading 1"].font.size.pt == 15
    assert str(exported.styles["Heading 1"].font.color.rgb) == "000000"
    assert exported.paragraphs[-1].text == "引用证据索引"

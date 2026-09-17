from __future__ import annotations

from io import BytesIO

from docx import Document as WordDocument
from docx.oxml.ns import qn

from app.services.markdown_docx import (
    configure_document_styles,
    markdown_to_docx_bytes,
    set_run_font,
)

_MARKDOWN = (
    "# 顶层标题被忽略\n\n"
    "## 1. 报告摘要\n\n"
    "**结论：**报告期已说明。\n\n"
    "### 1.1 子项\n\n"
    "- 第一条要点\n"
    "- 第二条要点\n\n"
    "1. 有序第一项\n\n"
    "## 2. 数据表\n\n"
    "| 序号 | 检查项 | 结论 |\n"
    "| --- | --- | --- |\n"
    "| 1 | 报告期 | 满足 |\n"
    "| 2 | 排放因子 | 部分满足 |\n"
)


def test_markdown_renders_headings_lists_and_tables() -> None:
    stream = markdown_to_docx_bytes(_MARKDOWN)
    exported = WordDocument(BytesIO(stream.read()))

    text = "\n".join(paragraph.text for paragraph in exported.paragraphs)
    # 一级标题不进正文（由封面承载），二级降为 Heading 1，三级降为 Heading 2
    assert "顶层标题被忽略" not in text
    assert "1. 报告摘要" in text
    assert "1.1 子项" in text
    assert "结论：报告期已说明。" in text

    # 表格被解析为真实表格：表头行 + 两行数据
    assert len(exported.tables) == 1
    table = exported.tables[0]
    assert len(table.rows) == 3
    assert [cell.text for cell in table.rows[0].cells] == ["序号", "检查项", "结论"]
    assert table.cell(1, 1).text == "报告期"
    assert table.cell(2, 2).text == "部分满足"

    styles = [paragraph.style.name for paragraph in exported.paragraphs]
    assert "Heading 1" in styles
    assert "Heading 2" in styles
    assert "List Bullet" in styles
    assert "List Number" in styles


def test_markdown_renders_bold_run_as_bold() -> None:
    stream = markdown_to_docx_bytes("普通文字与**加粗结论**并列。")
    exported = WordDocument(BytesIO(stream.read()))
    paragraph = next(item for item in exported.paragraphs if item.text.startswith("普通文字"))
    bold_runs = [run.text for run in paragraph.runs if run.bold]
    assert bold_runs == ["加粗结论"]


def test_configure_document_styles_applies_public_document_typography() -> None:
    document = WordDocument()
    configure_document_styles(document)

    assert document.styles["Normal"].font.size.pt == 12
    assert document.styles["Heading 1"].font.size.pt == 15
    assert document.styles["Heading 2"].font.size.pt == 14
    assert str(document.styles["Heading 1"].font.color.rgb) == "000000"


def test_set_run_font_switches_east_asian_font_family() -> None:
    document = WordDocument()
    run = document.add_paragraph().add_run("测试")
    set_run_font(run, size=14, bold=True, east_asia="仿宋_GB2312")

    assert run.font.size.pt == 14
    assert run.bold is True
    assert run._element.rPr.rFonts.get(qn("w:eastAsia")) == "仿宋_GB2312"
    assert run._element.rPr.rFonts.get(qn("w:ascii")) == "Times New Roman"


def test_markdown_to_docx_bytes_enables_field_update() -> None:
    stream = markdown_to_docx_bytes("## 章节\n\n正文。")
    exported = WordDocument(BytesIO(stream.read()))

    update_fields = exported.settings._element.find(qn("w:updateFields"))
    assert update_fields is not None
    assert update_fields.get(qn("w:val")) == "true"


def test_table_separator_without_pipe_row_is_not_a_table() -> None:
    stream = markdown_to_docx_bytes("## 章节\n\n| --- | --- |\n")
    exported = WordDocument(BytesIO(stream.read()))
    assert exported.tables == []

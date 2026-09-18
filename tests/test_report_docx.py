"""Word 导出：表格必须是表格，封面不能出现内部编号。

这一组用例钉的是「ReportIR 直渲染 Word」这条路线上的回归点，每一条都真的坏过：

* 表格块少一个列名（模型写了 ``headers`` 而不是 ``columns``），整张表就在 Word 里
  退化成一段带竖线的文字，看上去像漏排；
* 证据标注挂在表格最后一行上，Markdown 把它当成一个单元格，两列的表变成三列；
* 封面写着「R1 报告」和「文档版本 vNone」——R1 是内部编号，直传材料没有文档版本；
* 分项读不出来的块被静默丢掉，纸面上少一段内容而没有任何提示。
"""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document as WordDocument

from app.domain.models import ReportRun
from app.services.report_docx import render_report_docx
from app.services.report_templates import ReportTemplate, ReportTemplateRegistry

WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
RUN_ID = "fe7d0088-1050-4333-9b27-72e966475ff3"


def _reporting_root() -> Path:
    return Path(__file__).resolve().parents[1] / "reporting"


def _template() -> ReportTemplate:
    return ReportTemplateRegistry(_reporting_root()).get("R1")


def _run(**overrides: Any) -> ReportRun:
    values: dict[str, Any] = {
        "id": RUN_ID,
        "user_id": 11,
        "report_type": "R1",
        "template_id": "r1-product-carbon-footprint",
        "template_version": "1.0.0",
        "mode": "GENERATE",
        "language": "zh-CN",
        "output_formats": ["ONLINE", "DOCX"],
        "llm_config_id": 5,
        "llm_snapshot_version": 2,
        "state": "SUCCEEDED",
        "stage": "COMPLETED",
        "input_hash": "a" * 64,
        "created_at": datetime(2026, 9, 18, 12, 48, 8, tzinfo=UTC),
        "finished_at": datetime(2026, 9, 18, 13, 12, 3, tzinfo=UTC),
    }
    values.update(overrides)
    return ReportRun(**values)


def _ir(blocks: list[dict[str, Any]], **overrides: Any) -> dict[str, Any]:
    report_ir: dict[str, Any] = {
        "schema_version": 1,
        "meta": {
            "run_id": RUN_ID,
            "report_type": "R1",
            "template_id": "r1-product-carbon-footprint",
            "template_version": "1.0.0",
            "document_id": None,
            "document_version": None,
            "source_filename": "伊顿碳足迹报告.docx",
            "language": "zh-CN",
        },
        "evidence": [],
        "field_ledger": [],
        "sections": [{"section_id": "results", "title": "核查结果", "blocks": blocks}],
        "calculations": [],
        "warnings": [],
        "limitations": [],
        "render_profile": "zh-report-v1",
    }
    report_ir.update(overrides)
    return report_ir


def _document(blocks: list[dict[str, Any]], **overrides: Any):
    payload = render_report_docx(
        run=_run(**overrides.pop("run", {})),
        template=_template(),
        report_ir=_ir(blocks, **overrides),
    )
    assert payload[:2] == b"PK"  # zip 容器
    return WordDocument(BytesIO(payload))


def _body_text(document: WordDocument) -> str:
    """正文纯文本：段落 + 表格单元格。指标卡与图表的数据都在表格里。"""
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _table_block(data: dict[str, Any], **extra: Any) -> dict[str, Any]:
    return {
        "type": "table",
        "text": "系统边界与核算基准",
        "data": data,
        "evidence_ids": [],
        **extra,
    }


def test_table_block_becomes_a_real_word_table() -> None:
    document = _document(
        [
            _table_block(
                {
                    "columns": ["项目", "内容", "证据"],
                    "rows": [
                        ["系统边界", "摇篮到大门", "用户确认"],
                        ["功能单位", "1 P单元", "文档"],
                    ],
                }
            )
        ]
    )

    assert len(document.tables) == 1
    table = document.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == ["项目", "内容", "证据"]
    assert [cell.text for cell in table.rows[1].cells] == ["系统边界", "摇篮到大门", "用户确认"]
    # 带竖线的段落意味着表格没被识别出来——那正是这次要修的毛病
    assert "|" not in _body_text(document)
    # 表格标题（块的 text）也要落到纸面上，不能因为它是标题就消失
    assert "系统边界与核算基准" in _body_text(document)


def test_table_with_legacy_headers_key_still_becomes_a_table() -> None:
    """模型写过 headers 而不是 columns，存量 IR 也必须出表格。"""
    document = _document(
        [_table_block({"headers": ["项目", "值"], "rows": [["功能单位", "1 吨"]]})]
    )

    assert len(document.tables) == 1
    assert [cell.text for cell in document.tables[0].rows[0].cells] == ["项目", "值"]
    assert "|" not in _body_text(document)


def test_table_without_columns_anywhere_is_still_a_table() -> None:
    """一个列名都没有时补空表头，也不能退回成一段带竖线的文字。"""
    document = _document(
        [
            {
                "type": "table",
                "text": "阶段数据",
                "data": {"rows": [["原材料", "392.04"], ["运输", "15.16"]]},
                "evidence_ids": [],
            }
        ]
    )

    assert len(document.tables) == 1
    assert document.tables[0].rows[1].cells[0].text == "原材料"
    assert "|" not in _body_text(document)


def test_evidence_citation_lands_outside_the_table() -> None:
    """证据标注不能进单元格：进去就会多撑出一列，两列的表看着像三列。"""
    document = _document(
        [
            _table_block(
                {"columns": ["项目", "值"], "rows": [["功能单位", "1 吨"]]},
                evidence_ids=["E-DOC-FU", "E-DOC-GWP"],
            )
        ]
    )

    table = document.tables[0]
    assert len(table.columns) == 2
    assert all("证据" not in cell.text for row in table.rows for cell in row.cells)
    assert "（证据：E-DOC-FU、E-DOC-GWP）" in _body_text(document)


def test_cover_shows_the_report_name_and_no_internal_identifiers() -> None:
    document = _document([{"type": "paragraph", "text": "正文。", "evidence_ids": []}])
    text = _body_text(document)

    assert "产品碳足迹评价报告" in text
    assert "生成日期 2026-09-18" in text
    # 任务 UUID、模板 slug、以及直传材料才会出现的 vNone，都不该出现在正文里
    assert RUN_ID not in text
    assert "r1-product-carbon-footprint" not in text
    assert "vNone" not in text
    assert "R1" not in text


def test_footer_has_page_number_and_run_number() -> None:
    document = _document([{"type": "paragraph", "text": "正文。", "evidence_ids": []}])
    footer = document.sections[0].footer.paragraphs[0]

    assert "PAGE" in footer._p.xml  # 页码是域，Word 打开时刷新
    assert "第 " in footer.text and footer.text.strip().endswith("页")
    assert "fe7d0088" in footer.text


def test_page_break_block_emits_a_real_break() -> None:
    document = _document(
        [
            {"type": "paragraph", "text": "第一页。", "evidence_ids": []},
            {"type": "page_break", "evidence_ids": []},
            {"type": "paragraph", "text": "第二页。", "evidence_ids": []},
        ]
    )

    breaks = document.element.body.findall(f".//{{{WORD_NS}}}br")
    assert any(node.get(f"{{{WORD_NS}}}type") == "page" for node in breaks)


def test_metric_cards_with_cards_alias_are_rendered() -> None:
    """模型写成 cards 的指标卡曾经整块消失，纸面上只剩一个空白。"""
    document = _document(
        [
            {
                "type": "metric_cards",
                "text": "关键指标",
                "data": {
                    "cards": [
                        {"label": "产品名称", "value": "MCB微型断路器"},
                        {"label": "碳足迹总值", "value": "0.44131 kgCO2e"},
                    ]
                },
                "evidence_ids": [],
            }
        ]
    )
    text = _body_text(document)

    assert "关键指标" in text
    assert "MCB微型断路器" in text
    assert "0.44131 kgCO2e" in text


def test_bar_chart_with_categories_and_values_renders_bars_and_data_table() -> None:
    """图表块在 Word 里落成条形 + 等价数据表；模型写成 categories/values 也要认。"""
    document = _document(
        [
            {
                "type": "bar_chart",
                "text": "生命周期各阶段排放",
                "data": {
                    "unit": "kgCO2e",
                    "categories": ["原料获取", "制造阶段", "运输分销"],
                    "values": [0.39204, 0.03411, 0.01516],
                },
                "evidence_ids": ["E-DOC-RESULT"],
            }
        ]
    )
    text = _body_text(document)
    xml = document.element.body.xml

    assert "生命周期各阶段排放" in text
    # 条形是靠单元格底色画的：没有底色就只剩一张表，看不出长短
    assert "2B8A63" in xml and "EDEAE1" in xml
    # 每张图都附一份等价数据表，数字能复制、能核对
    assert "分项" in text and "数值（kgCO2e）" in text
    assert "0.39204" in text


def test_unreadable_block_falls_back_to_raw_instead_of_disappearing() -> None:
    """分项读不出来时按原样落到纸面。少一段内容而没有任何痕迹，是最难发现的那种错。"""
    document = _document(
        [
            {
                "type": "metric_cards",
                "text": "关键指标",
                "data": {"unknown_shape": [{"名字": "产品名称", "取值": "MCB"}]},
                "evidence_ids": [],
            }
        ]
    )
    text = _body_text(document)

    assert "unknown_shape" in text
    assert "MCB" in text


def test_empty_section_keeps_no_heading() -> None:
    """整节都排不出内容时不留一个孤零零的标题。"""
    document = _document([{"type": "paragraph", "text": "", "evidence_ids": []}])

    assert "核查结果" not in _body_text(document)

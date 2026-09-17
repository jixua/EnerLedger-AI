"""报告产物：把通过校验的 ReportIR 渲染成可下载的 Markdown / DOCX 并落盘。

设计上与「文档分析报告」保持一致：渲染产物写入对象存储，`report_artifact` 表登记
（run_id + artifact_type 唯一），前端凭接口拿到的 artifact 条目下载。渲染是纯函数，
不触发任何模型请求；落盘失败不会影响报告本身的生成结果（报告已在库里）。
"""

from __future__ import annotations

import asyncio
import contextlib
import hashlib
import json
import os
import re
import tempfile
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.domain.models import ReportArtifact, ReportRun
from app.rag.config import settings
from app.rag.observability.logging import logger
from app.rag.services.storage.base import BaseObjectStorage
from app.rag.services.storage.factory import StorageFactory
from app.services.document_analysis_docx import (
    _add_inline,
    _add_markdown,
    _configure_styles,
    _set_run_font,
)
from app.services.report_templates import ReportTemplate

RENDERER_VERSION = "report-ir/1"

# 除在线渲染外允许落盘的格式。
DOWNLOADABLE_FORMATS = ("MARKDOWN", "DOCX")
_CONTENT_TYPES = {
    "MARKDOWN": "text/markdown; charset=utf-8",
    "DOCX": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
_EXTENSIONS = {"MARKDOWN": "md", "DOCX": "docx"}


class ReportArtifactError(RuntimeError):
    code = "REPORT_ARTIFACT_FAILED"


def render_report_markdown(
    *, run: ReportRun, template: ReportTemplate, report_ir: dict[str, Any]
) -> str:
    """把 ReportIR 渲染成 Markdown（同时作为 DOCX 的中间表示）。

    不使用引用块（``>``）语法：Word 转换只识别标题/表格/列表/段落，
    提示类内容用「【提示】」前缀表达，两种产物都成立。
    """
    lines: list[str] = [f"# {run.report_type} 报告 · {template.definition.get('name') or run.template_id}"]
    lines.append("")
    for section in report_ir.get("sections") or []:
        lines.append(f"## {section.get('title')}")
        lines.append("")
        for block in section.get("blocks") or []:
            rendered = _render_block(block)
            if rendered:
                lines.append(rendered)
                lines.append("")

    ledger = report_ir.get("field_ledger") or []
    if ledger:
        lines += ["## 附表一：字段台账", "", "| 字段 | 状态 | 值 | 单位 | 证据 |", "| --- | --- | --- | --- | --- |"]
        lines += [
            "| {field} | {status} | {value} | {unit} | {evidence} |".format(
                field=item.get("field_id"),
                status=item.get("status"),
                value=_cell(item.get("value")),
                unit=_cell(item.get("unit")),
                evidence=", ".join(item.get("evidence_ids") or []),
            )
            for item in ledger
        ]
        lines.append("")

    evidence = report_ir.get("evidence") or []
    if evidence:
        lines += ["## 附表二：证据台账", "", "| 证据 | 来源 | 位置 | 摘录 |", "| --- | --- | --- | --- |"]
        lines += [
            "| {evidence_id} | {source_type} | {location} | {excerpt} |".format(
                evidence_id=item.get("evidence_id"),
                source_type=item.get("source_type"),
                location=_cell(_evidence_location(item)),
                excerpt=_cell(_trim(item.get("excerpt"), 120)),
            )
            for item in evidence
        ]
        lines.append("")

    calculations = report_ir.get("calculations") or []
    if calculations:
        lines += ["## 附表三：计算过程", ""]
        for item in calculations:
            lines.append(
                "- `{formula}` v{version}：{field} = {value} {unit}（输入：{inputs}）".format(
                    formula=item.get("formula_id"),
                    version=item.get("formula_version"),
                    field=item.get("output_field_id"),
                    value=item.get("value"),
                    unit=item.get("unit") or "",
                    inputs=json.dumps(item.get("inputs"), ensure_ascii=False),
                )
            )
        lines.append("")

    for key, title in (("limitations", "限制声明"), ("warnings", "警告")):
        items = report_ir.get(key) or []
        if items:
            lines += [f"## 附表四：{title}" if key == "limitations" else f"## 附表五：{title}", ""]
            lines += [f"- {item}" for item in items]
            lines.append("")
    return "\n".join(lines).strip() + "\n"


def _render_block(block: dict[str, Any]) -> str:
    block_type = block.get("type")
    data = block.get("data") if isinstance(block.get("data"), dict) else {}
    evidence = block.get("evidence_ids") or []
    suffix = f"（证据：{', '.join(evidence)}）" if evidence else ""
    lines: list[str] = []

    if block_type in ("paragraph", "heading", "signature_block", "source_note"):
        text = str(block.get("text") or "").strip()
        if text:
            lines.append(f"{text}{suffix}")
    elif block_type == "callout":
        text = str(block.get("text") or "").strip()
        if text:
            lines.append(f"**【提示】** {text}{suffix}")
    elif block_type == "list":
        items = data.get("items") or []
        texts = [str(item.get("text") if isinstance(item, dict) else item) for item in items]
        if not texts and block.get("text"):
            texts = [str(block["text"])]
        lines += [f"- {text}" for text in texts if text]
    elif block_type == "metric_cards":
        items = data.get("items") or []
        if items:
            lines += ["| 指标 | 值 |", "| --- | --- |"]
            lines += [f"| {_cell(item.get('label'))} | {_cell(item.get('value'))} |" for item in items]
    elif block_type in ("bar_chart", "donut_chart"):
        series = data.get("series") or []
        if series:
            unit = data.get("unit") or ""
            lines += [f"| 分项 | 数值（{unit}） |", "| --- | --- |"]
            lines += [f"| {_cell(item.get('label'))} | {_cell(item.get('value'))} |" for item in series]
    elif block_type == "table":
        rows = data.get("rows") or []
        columns = data.get("columns")
        if rows:
            if isinstance(rows[0], dict):
                columns = columns or list(rows[0].keys())
                lines += ["| " + " | ".join(_cell(name) for name in columns) + " |"]
                lines += ["|" + " --- |" * len(columns)]
                lines += [
                    "| " + " | ".join(_cell(row.get(name)) for name in columns) + " |" for row in rows
                ]
            else:
                if columns:
                    lines += ["| " + " | ".join(_cell(name) for name in columns) + " |"]
                    lines += ["|" + " --- |" * len(columns)]
                lines += ["| " + " | ".join(_cell(cell) for cell in row) + " |" for row in rows]
        elif block.get("text"):
            lines.append(str(block["text"]))
    elif block_type == "page_break":
        lines.append("")
    else:
        payload = json.dumps(data or block.get("text"), ensure_ascii=False)
        lines.append(f"```\n{_trim(payload, 1200)}\n```")

    if suffix and lines and not any(suffix in line for line in lines):
        lines[-1] = f"{lines[-1]}{suffix}"
    return "\n".join(lines)


def _cell(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).replace("|", "／").strip()


def _trim(value: Any, limit: int) -> str:
    text = "" if value is None else str(value)
    text = re.sub(r"\s+", " ", text).strip()
    return text[:limit]


def _evidence_location(item: dict[str, Any]) -> str:
    if item.get("chunk_id"):
        return str(item["chunk_id"])
    if item.get("reference_uri"):
        return str(item["reference_uri"])
    if item.get("page"):
        return f"第 {item['page']} 页"
    return ""


def build_report_docx(*, run: ReportRun, template: ReportTemplate, markdown: str) -> bytes:
    """把渲染好的 Markdown 转成 Word（复用文档分析报告的排版原子）。"""
    output = WordDocument()
    _configure_styles(output)
    output.core_properties.title = f"{run.report_type} 报告"
    output.core_properties.subject = template.definition.get("name") or run.template_id

    section = output.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(2.5)

    title = output.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(title.add_run(f"{run.report_type} 报告"), size=22, bold=True)
    subtitle = output.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(subtitle.add_run(template.definition.get("name") or run.template_id), size=13)

    meta = output.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(
        meta.add_run(
            f"任务 {run.id}　文档版本 v{run.document_version}　"
            f"模板 {run.template_id} v{run.template_version}"
        ),
        size=9,
    )
    output.add_paragraph()

    _add_markdown(output, markdown)

    stream = BytesIO()
    output.save(stream)
    return stream.getvalue()


async def read_report_artifact(
    artifact: ReportArtifact,
    *,
    max_bytes: int = 32 * 1024 * 1024,
    storage: BaseObjectStorage | None = None,
) -> bytes:
    """读取产物内容（报告产物都很小，走临时文件 + 上限保护，避免违反存储层的内存约定）。"""
    object_storage = storage or StorageFactory.get_storage()
    temp_path: Path | None = None
    try:
        descriptor, temp_name = tempfile.mkstemp(prefix="report-artifact-", suffix=".tmp")
        temp_path = Path(temp_name)
        os.close(descriptor)
        await asyncio.to_thread(
            object_storage.download_to_path, artifact.bucket, artifact.object_key, temp_path
        )
        if temp_path.stat().st_size > max_bytes:
            raise ReportArtifactError("报告产物超过读取上限")
        if temp_path.stat().st_size != int(artifact.size_bytes or 0):
            raise ReportArtifactError("报告产物大小与登记不一致，请重新生成")
        return await asyncio.to_thread(temp_path.read_bytes)
    except ReportArtifactError:
        raise
    except Exception as exc:
        raise ReportArtifactError(f"报告产物读取失败：{exc}") from exc
    finally:
        if temp_path is not None:
            with contextlib.suppress(OSError):
                temp_path.unlink()


def artifact_object_key(*, run: ReportRun, artifact_type: str, digest: str) -> str:
    return (
        f"reports/{int(run.user_id)}/{int(run.dataset_id)}/{int(run.document_id)}/"
        f"{run.id}/report-{digest[:12]}.{_EXTENSIONS[artifact_type]}"
    )


def artifact_download_name(*, run: ReportRun, document_filename: str | None, artifact_type: str) -> str:
    stem = re.split(r"[\\/]", str(document_filename or ""))[-1]
    stem = re.sub(r"\.[^.]+$", "", stem).strip() or f"报告-{run.id[:8]}"
    stem = re.sub(r'[\\/:*?"<>|]+', "_", stem)[:60]
    return f"{stem}-{run.report_type}报告.{_EXTENSIONS[artifact_type]}"


async def persist_report_artifacts(
    db: AsyncSession,
    *,
    run: ReportRun,
    template: ReportTemplate,
    report_ir: dict[str, Any],
    storage: BaseObjectStorage | None = None,
) -> list[ReportArtifact]:
    """按 run 声明的格式渲染并落盘产物；重复执行会先清掉本 run 的旧产物。"""
    requested = [fmt for fmt in DOWNLOADABLE_FORMATS if fmt in set(run.output_formats or [])]
    if not requested:
        return []

    object_storage = storage or StorageFactory.get_storage()
    bucket = settings.MINIO_PRIVATE_BUCKET
    markdown = render_report_markdown(run=run, template=template, report_ir=report_ir)
    payloads: list[tuple[str, bytes]] = []
    if "MARKDOWN" in requested:
        payloads.append(("MARKDOWN", markdown.encode("utf-8")))
    if "DOCX" in requested:
        payloads.append(("DOCX", build_report_docx(run=run, template=template, markdown=markdown)))

    existing = (
        await db.scalars(select(ReportArtifact).where(ReportArtifact.run_id == str(run.id)))
    ).all()
    for artifact in existing:
        try:
            object_storage.remove_object(artifact.bucket, artifact.object_key)
        except Exception as exc:  # 对象可能已不存在，不影响重写
            logger.bind(
                event="report_artifact_object_remove_failed",
                run_id=str(run.id),
                artifact_type=artifact.artifact_type,
                error=str(exc)[:200],
            ).warning("清理旧报告产物失败")
        await db.delete(artifact)

    created: list[ReportArtifact] = []
    for artifact_type, data in payloads:
        digest = hashlib.sha256(data).hexdigest()
        object_key = artifact_object_key(run=run, artifact_type=artifact_type, digest=digest)
        try:
            object_storage.upload_bytes(bucket, object_key, data, _CONTENT_TYPES[artifact_type])
        except Exception as exc:
            raise ReportArtifactError(f"{artifact_type} 产物上传失败：{exc}") from exc
        artifact = ReportArtifact(
            run_id=str(run.id),
            artifact_type=artifact_type,
            bucket=bucket,
            object_key=object_key,
            content_type=_CONTENT_TYPES[artifact_type],
            content_hash=digest,
            size_bytes=len(data),
            renderer_version=RENDERER_VERSION,
        )
        db.add(artifact)
        created.append(artifact)

    await db.commit()
    logger.bind(
        event="report_artifacts_persisted",
        run_id=str(run.id),
        formats=[artifact.artifact_type for artifact in created],
        bytes=sum(artifact.size_bytes for artifact in created),
    ).info("报告产物已落盘")
    return created

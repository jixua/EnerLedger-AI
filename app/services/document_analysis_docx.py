"""将已持久化的企业文档 Markdown 分析报告导出为 DOCX。"""

from __future__ import annotations

import re
from io import BytesIO

from docx import Document as WordDocument
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.domain.models import Document
from app.services.document_analysis import AnalysisSource, DocumentAnalysisResult

_TABLE_SEPARATOR = re.compile(r"^\s*:?-{3,}:?\s*$")
_INLINE_TOKEN = re.compile(r"(\*\*.+?\*\*|\[\u6587\u6863\u7247\u6bb5\d+\])")
_BODY_FONT = "\u5b8b\u4f53"
_HEADING_FONT = "\u9ed1\u4f53"
_COVER_FONT = "\u5fae\u8f6f\u96c5\u9ed1"
_LATIN_FONT = "Times New Roman"
_TABLE_WIDTH_DXA = 8220


def _set_run_font(
    run,
    *,
    size: float = 12,
    bold: bool | None = None,
    east_asia: str = _BODY_FONT,
) -> None:
    run.font.name = _LATIN_FONT
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), _LATIN_FONT)
    fonts.set(qn("w:hAnsi"), _LATIN_FONT)
    fonts.set(qn("w:cs"), _LATIN_FONT)
    fonts.set(qn("w:eastAsia"), east_asia)
    if bold is not None:
        run.bold = bold


def _set_cell_margins(
    cell,
    *,
    top: int = 90,
    start: int = 110,
    bottom: int = 90,
    end: int = 110,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    _set_run_font(run, size=10.5)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))


def _set_style_fonts(style, *, east_asia: str, size: float, bold: bool | None) -> None:
    style.font.name = _LATIN_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    fonts = style._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), _LATIN_FONT)
    fonts.set(qn("w:hAnsi"), _LATIN_FONT)
    fonts.set(qn("w:cs"), _LATIN_FONT)
    fonts.set(qn("w:eastAsia"), east_asia)


def _configure_styles(document: WordDocument) -> None:
    normal = document.styles["Normal"]
    _set_style_fonts(normal, east_asia=_BODY_FONT, size=12, bold=None)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for style_name, size, font, line_spacing, before, after in (
        ("Heading 1", 15, _HEADING_FONT, 2.0, 1, 1),
        ("Heading 2", 14, _HEADING_FONT, 1.25, 6, 6),
        ("Heading 3", 12, _BODY_FONT, 1.25, 12, 0),
    ):
        style = document.styles[style_name]
        _set_style_fonts(style, east_asia=font, size=size, bold=True)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing = line_spacing
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def _add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in _INLINE_TOKEN.finditer(text):
        if match.start() > cursor:
            _set_run_font(paragraph.add_run(text[cursor : match.start()]))
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run, bold=True)
        else:
            run = paragraph.add_run(token)
            _set_run_font(run, size=10.5)
        cursor = match.end()
    if cursor < len(text):
        _set_run_font(paragraph.add_run(text[cursor:]))


def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_table(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines) or "|" not in lines[index]:
        return False
    separators = _split_table_row(lines[index + 1])
    return bool(separators) and all(_TABLE_SEPARATOR.match(cell) for cell in separators)


def _add_table(document: WordDocument, rows: list[list[str]]) -> None:
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_pr = table._tbl.tblPr
    table_width = table_pr.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(_TABLE_WIDTH_DXA))
    table_width.set(qn("w:type"), "dxa")
    table_indent = OxmlElement("w:tblInd")
    table_indent.set(qn("w:w"), "0")
    table_indent.set(qn("w:type"), "dxa")
    table_pr.append(table_indent)

    max_lengths = [
        max(len(row[index]) if index < len(row) else 0 for row in rows)
        for index in range(column_count)
    ]
    weights = [max(4, min(length, 28)) for length in max_lengths]
    weight_sum = sum(weights)
    widths = [max(760, round(_TABLE_WIDTH_DXA * weight / weight_sum)) for weight in weights]
    widths[-1] += _TABLE_WIDTH_DXA - sum(widths)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        for column_index in range(column_count):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            tc_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_width.set(qn("w:w"), str(widths[column_index]))
            tc_width.set(qn("w:type"), "dxa")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.space_after = Pt(0)
            _add_inline(paragraph, values[column_index] if column_index < len(values) else "")
            for run in paragraph.runs:
                _set_run_font(run, size=10.5, bold=row_index == 0)
        if row_index == 0:
            _repeat_table_header(row)


def _add_markdown(document: WordDocument, markdown: str) -> None:
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line.strip():
            index += 1
            continue
        if _is_table(lines, index):
            rows = [_split_table_row(line)]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(_split_table_row(lines[index]))
                index += 1
            _add_table(document, rows)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            if level > 1:
                document.add_paragraph(
                    heading.group(2).strip(),
                    style=f"Heading {min(level - 1, 3)}",
                )
            index += 1
            continue
        bullet = re.match(r"^\s*[-*+]\s+(.+)$", line)
        numbered = re.match(r"^\s*\d+[.)]\s+(.+)$", line)
        if bullet or numbered:
            paragraph = document.add_paragraph(style="List Bullet" if bullet else "List Number")
            paragraph.paragraph_format.line_spacing = 1.25
            paragraph.paragraph_format.space_after = Pt(0)
            _add_inline(paragraph, (bullet or numbered).group(1))
            index += 1
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(24)
        paragraph.paragraph_format.line_spacing = 1.25
        _add_inline(paragraph, line.strip())
        index += 1


def _source_location(source: AnalysisSource) -> str:
    if source.page is not None:
        return f"\u7b2c {source.page} \u9875"
    if source.page_range:
        return f"\u7b2c {source.page_range['start']}-{source.page_range['end']} \u9875"
    return "\u9875\u7801\u672a\u8bb0\u5f55"


def build_document_analysis_docx(*, document: Document, result: DocumentAnalysisResult) -> bytes:
    """根据已持久化的报告生成 Word 文件，不触发新的模型请求。"""

    output = WordDocument()
    _configure_styles(output)
    output.core_properties.title = "企业文档分析报告"
    output.core_properties.subject = "大模型辅助企业文档证据分析"
    section = output.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(3.0)
    section.top_margin = Cm(3.5)
    section.bottom_margin = Cm(3.0)
    section.header_distance = Cm(2.8)
    section.footer_distance = Cm(2.0)

    section.header.paragraphs[0].clear()
    _add_page_number(section.footer.paragraphs[0])

    report_number = output.add_paragraph()
    report_number.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_font(
        report_number.add_run(f"分析报告编号：AI{document.id:06d}-V{document.version}"),
        size=11,
    )
    for _ in range(3):
        output.add_paragraph()

    source_title = re.sub(r"\.[^.]+$", "", document.filename).strip() or "企业文档"
    for text in (source_title, "企业文档", "分析报告"):
        title = output.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.line_spacing = 1.25
        _set_run_font(title.add_run(text), size=24, east_asia=_COVER_FONT)

    for _ in range(5):
        output.add_paragraph()
    for label, value in (
        ("\u6e90\u6587\u4ef6", document.filename),
        ("\u6587\u6863\u7248\u672c", str(document.version)),
        (
            "\u751f\u6210\u65f6\u95f4",
            result.generated_at.astimezone().strftime("%Y-%m-%d %H:%M:%S"),
        ),
        ("\u5206\u6790\u6a21\u578b", result.model_name),
    ):
        paragraph = output.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Cm(1.0)
        paragraph.paragraph_format.line_spacing = 1.0
        _set_run_font(
            paragraph.add_run(f"{label}\uff1a"),
            size=14,
            bold=True,
            east_asia="仿宋_GB2312",
        )
        _set_run_font(
            paragraph.add_run(str(value)),
            size=14,
            east_asia="仿宋_GB2312",
        )
    output.add_paragraph()
    disclaimer = output.add_paragraph()
    disclaimer.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    disclaimer.paragraph_format.first_line_indent = Pt(21)
    disclaimer.paragraph_format.line_spacing = 1.25
    run = disclaimer.add_run(
        "\u58f0\u660e：\u672c\u62a5\u544a\u7531\u5927\u6a21\u578b\u57fa\u4e8e\u5f53\u524d\u6587\u6863\u5185\u5bb9\u8f85\u52a9\u751f\u6210，"
        "\u4e0d\u66ff\u4ee3\u6b63\u5f0f\u7b2c\u4e09\u65b9\u6838\u67e5\u3001\u5ba1\u8ba1\u3001\u8ba4\u8bc1\u6216\u6cd5\u89c4\u610f\u89c1。"
    )
    _set_run_font(run, size=10.5)

    output.add_section(WD_SECTION.NEW_PAGE)
    content_section = output.sections[-1]
    content_section.header.is_linked_to_previous = True
    content_section.footer.is_linked_to_previous = True
    _add_markdown(output, result.markdown)

    output.add_paragraph("\u5f15\u7528\u8bc1\u636e\u7d22\u5f15", style="Heading 1")
    if result.sources:
        source_rows = [["\u5f15\u7528", "\u4f4d\u7f6e", "\u5206\u7247", "\u8bc1\u636e\u6458\u8981"]]
        source_rows.extend(
            [
                f"[\u6587\u6863\u7247\u6bb5{source.citation_index}]",
                _source_location(source),
                str(source.chunk_index + 1),
                source.excerpt or "\u672a\u63d0\u4f9b\u6458\u8981",
            ]
            for source in result.sources
        )
        _add_table(output, source_rows)
    else:
        output.add_paragraph("\u672c\u6b21\u5206\u6790\u672a\u5f62\u6210\u53ef\u6620\u5c04\u7684\u6587\u6863\u7247\u6bb5\u5f15\u7528。")

    settings = output.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    stream = BytesIO()
    output.save(stream)
    return stream.getvalue()

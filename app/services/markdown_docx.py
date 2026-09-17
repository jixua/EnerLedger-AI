"""把 Markdown 渲染为符合中文公文样式的 DOCX。

报告产物（``app/services/report_artifacts.py``）与后续其他导出场景共用这里的
字体、表格和标题实现，避免各自维护一套排版规则。
"""

from __future__ import annotations

import re
from io import BytesIO

from docx import Document as WordDocument
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

_TABLE_SEPARATOR = re.compile(r"^\s*:?-{3,}:?\s*$")
_INLINE_TOKEN = re.compile(r"(\*\*.+?\*\*)")
BODY_FONT = "宋体"
HEADING_FONT = "黑体"
COVER_FONT = "微软雅黑"
LATIN_FONT = "Times New Roman"
_TABLE_WIDTH_DXA = 8220


def set_run_font(
    run,
    *,
    size: float = 12,
    bold: bool | None = None,
    east_asia: str = BODY_FONT,
) -> None:
    run.font.name = LATIN_FONT
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:cs"), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), east_asia)
    if bold is not None:
        run.bold = bold


def _set_cell_margins(
    cell,
    *,
    top: int = 90,
    start: int = 110,
    bottom: int = 90,
    end: int = 110,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=10.5)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))


def enable_field_update(document: WordDocument) -> None:
    """让 Word 打开文档时自动刷新页码等域。"""

    settings_element = document.settings._element
    update_fields = settings_element.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings_element.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def _set_style_fonts(style, *, east_asia: str, size: float, bold: bool | None) -> None:
    style.font.name = LATIN_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    fonts = style._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:cs"), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), east_asia)


def configure_document_styles(document: WordDocument) -> None:
    normal = document.styles["Normal"]
    _set_style_fonts(normal, east_asia=BODY_FONT, size=12, bold=None)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for style_name, size, font, line_spacing, before, after in (
        ("Heading 1", 15, HEADING_FONT, 2.0, 1, 1),
        ("Heading 2", 14, HEADING_FONT, 1.25, 6, 6),
        ("Heading 3", 12, BODY_FONT, 1.25, 12, 0),
    ):
        style = document.styles[style_name]
        _set_style_fonts(style, east_asia=font, size=size, bold=True)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing = line_spacing
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def _add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in _INLINE_TOKEN.finditer(text):
        if match.start() > cursor:
            set_run_font(paragraph.add_run(text[cursor : match.start()]))
        set_run_font(paragraph.add_run(match.group(0)[2:-2]), bold=True)
        cursor = match.end()
    if cursor < len(text):
        set_run_font(paragraph.add_run(text[cursor:]))


def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_table(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines) or "|" not in lines[index]:
        return False
    separators = _split_table_row(lines[index + 1])
    return bool(separators) and all(_TABLE_SEPARATOR.match(cell) for cell in separators)


def add_table(document: WordDocument, rows: list[list[str]]) -> None:
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_pr = table._tbl.tblPr
    table_width = table_pr.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(_TABLE_WIDTH_DXA))
    table_width.set(qn("w:type"), "dxa")
    table_indent = OxmlElement("w:tblInd")
    table_indent.set(qn("w:w"), "0")
    table_indent.set(qn("w:type"), "dxa")
    table_pr.append(table_indent)

    max_lengths = [
        max(len(row[index]) if index < len(row) else 0 for row in rows)
        for index in range(column_count)
    ]
    weights = [max(4, min(length, 28)) for length in max_lengths]
    weight_sum = sum(weights)
    widths = [max(760, round(_TABLE_WIDTH_DXA * weight / weight_sum)) for weight in weights]
    widths[-1] += _TABLE_WIDTH_DXA - sum(widths)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        for column_index in range(column_count):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            tc_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_width.set(qn("w:w"), str(widths[column_index]))
            tc_width.set(qn("w:type"), "dxa")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.space_after = Pt(0)
            _add_inline(paragraph, values[column_index] if column_index < len(values) else "")
            for run in paragraph.runs:
                set_run_font(run, size=10.5, bold=row_index == 0)
        if row_index == 0:
            _repeat_table_header(row)


def add_markdown(document: WordDocument, markdown: str) -> None:
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line.strip():
            index += 1
            continue
        if _is_table(lines, index):
            rows = [_split_table_row(line)]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(_split_table_row(lines[index]))
                index += 1
            add_table(document, rows)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            if level > 1:
                document.add_paragraph(
                    heading.group(2).strip(),
                    style=f"Heading {min(level - 1, 3)}",
                )
            index += 1
            continue
        bullet = re.match(r"^\s*[-*+]\s+(.+)$", line)
        numbered = re.match(r"^\s*\d+[.)]\s+(.+)$", line)
        if bullet or numbered:
            paragraph = document.add_paragraph(style="List Bullet" if bullet else "List Number")
            paragraph.paragraph_format.line_spacing = 1.25
            paragraph.paragraph_format.space_after = Pt(0)
            _add_inline(paragraph, (bullet or numbered).group(1))
            index += 1
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(24)
        paragraph.paragraph_format.line_spacing = 1.25
        _add_inline(paragraph, line.strip())
        index += 1


def markdown_to_docx_bytes(markdown: str) -> BytesIO:
    """按公文样式把 Markdown 渲染成 DOCX 字节流。"""

    document = WordDocument()
    configure_document_styles(document)
    add_markdown(document, markdown)
    enable_field_update(document)
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream

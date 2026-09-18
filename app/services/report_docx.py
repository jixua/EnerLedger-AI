"""报告产物：把 ReportIR 渲染成 Word（DOCX）。

与 HTML 产物同一条路线——直接从 ReportIR 出发，不经过 Markdown 中间表示。原先的
路线是 IR → Markdown → 手写解析 → python-docx，问题全出在中间那一层：表格少一行
分隔符就退化成一段带竖线的文字、行内反引号原样留在正文里、分页符直接消失。这些
毛病要等文件生成之后有人打开 Word 才看得见，而且看不见的那部分内容没了也没人知道。

排版规则在 ``word_layout``（字体、样式、表格、页码），``report_blocks`` 负责把 IR
读成列／行／分项——同一个形状的判断在 HTML、Markdown、DOCX 三个渲染器之间共用。
这里只管「哪种块排成什么样子」。

图表不依赖绘图库：条形与占比条用表格单元格的底色画（底色连成一片就是一根条），
指标卡用无框表格排。渲染是纯函数，不触发模型请求。
"""

from __future__ import annotations

import json
import re
from io import BytesIO
from typing import Any

from docx import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.domain.models import ReportRun
from app.services.report_blocks import (
    FIELD_STATUS_LABELS,
    UNRESOLVED_STATUSES,
    block_data,
    block_unit,
    evidence_location,
    flatten_text,
    list_items,
    metric_items,
    series_items,
    table_columns,
    table_rows,
    to_number,
    trim_text,
)
from app.services.report_templates import ReportTemplate
from app.services.word_layout import (
    COVER_FONT,
    HEADING_FONT,
    MONO_FONT,
    add_page_number,
    column_weights,
    configure_document_styles,
    create_table,
    enable_field_update,
    repeat_table_header,
    set_cell_shading,
    set_page_setup,
    set_paragraph_left_border,
    set_paragraph_shading,
    set_run_font,
)

# 与 HTML 产物同一套配色（pages.css 的报告样式），两份导出件摆在一起才像一份东西。
_ACCENT = "0D5A46"
_SERIES = ("2B8A63", "4874B0", "CF6A35", "6D5BB5", "B08A20")
_POSITIVE = "CF6A35"
_NEGATIVE = "4874B0"
_TRACK = "EDEAE1"
_AXIS = "D6D1C4"
_HEADER_FILL = "EEF3F0"
_CALLOUT_FILL = "F5F3EC"
_MUTED = "6B7280"

_BODY_SIZE = 12
_TABLE_SIZE = 10.5
_NOTE_SIZE = 9
_COVER_TITLE_SIZE = 22

# 占比条最多画几段，超过就出数据表：色相要重复时读者无法把颜色对回分项。
_SHARE_SLOTS = 5
# 条形图的轨道格数。20 格 = 5% 一档，肉眼读得出长短；格子再多表格会变得密密麻麻。
_TRACK_CELLS = 20
_BAR_LABEL_WEIGHT = 24
_BAR_VALUE_WEIGHT = 20
_CJK = re.compile("[一-鿿]")


def _display(value: Any) -> str:
    """表格里的取值：空值统一成一个破折号，避免空格子被读成漏项。"""
    return flatten_text(value) or "—"


def _json_text(payload: Any) -> str:
    try:
        return json.dumps(payload, ensure_ascii=False, indent=2, default=str)
    except (TypeError, ValueError):
        return flatten_text(payload)


def _paragraph(
    document, *, size: float = _BODY_SIZE, indent: bool = False, style: str | None = None
):
    paragraph = document.add_paragraph(style=style)
    if indent:
        # 中文正文段首缩进两个字：不缩进在 Word 里读起来像列表
        paragraph.paragraph_format.first_line_indent = Pt(size * 2)
    paragraph.paragraph_format.space_after = Pt(0)
    return paragraph


def _add_paragraph(
    document, text: Any, *, indent: bool = True, style: str | None = None
) -> bool:
    """排一个段落，返回是否真的写下了内容——调用方据此决定要不要兜底。"""
    content = flatten_text(text)
    if not content:
        return False
    paragraph = _paragraph(document, indent=indent, style=style)
    set_run_font(paragraph.add_run(content))
    return True


def _add_callout(document, block: dict[str, Any]) -> bool:
    """提示块：左侧竖线 + 浅底色，与 HTML 产物的提示样式同一个意思。"""
    text = flatten_text(block.get("text"))
    if not text:
        return False
    paragraph = _paragraph(document, indent=False)
    paragraph.paragraph_format.left_indent = Pt(12)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    set_paragraph_shading(paragraph, _CALLOUT_FILL)
    set_paragraph_left_border(paragraph, _ACCENT)
    set_run_font(paragraph.add_run("提示　"), bold=True)
    set_run_font(paragraph.add_run(text))
    return True


def _add_caption(document, text: Any) -> None:
    """表格与图表的标题（块的 ``text``）。不排出来，模型写的这行标题就凭空消失了。"""
    content = flatten_text(text)
    if not content:
        return
    paragraph = _paragraph(document, size=_TABLE_SIZE, indent=False)
    paragraph.paragraph_format.space_before = Pt(6)
    set_run_font(paragraph.add_run(content), size=_TABLE_SIZE, bold=True, east_asia=HEADING_FONT)


def _add_note(document, text: str) -> None:
    paragraph = _paragraph(document, size=_NOTE_SIZE, indent=False)
    set_run_font(paragraph.add_run(text), size=_NOTE_SIZE, color=_MUTED)


def _add_evidence(document, evidence_ids: Any) -> None:
    """证据标注单独一行，不塞进任何单元格——塞进单元格会撑出一列新的空格子。"""
    if not isinstance(evidence_ids, list) or not evidence_ids:
        return
    ids = "、".join(item for item in (flatten_text(value) for value in evidence_ids) if item)
    if not ids:
        return
    paragraph = _paragraph(document, size=_NOTE_SIZE, indent=False)
    set_run_font(paragraph.add_run(f"（证据：{ids}）"), size=_NOTE_SIZE, color=_MUTED)


def _add_data_table(
    document,
    columns: list[str],
    rows: list[list[Any]],
    *,
    shading: str = _HEADER_FILL,
) -> None:
    """带表头的数据表：表头行加底色并在跨页时重复，数值列右对齐。"""
    if not rows:
        return
    widths = column_weights(columns, [[_display(cell) for cell in row] for row in rows])
    table = create_table(
        document, columns=len(columns), rows=len(rows) + 1, weights=widths
    )
    header = table.rows[0]
    for index, name in enumerate(columns):
        cell = header.cells[index]
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(paragraph.add_run(flatten_text(name)), size=_TABLE_SIZE, bold=True)
        set_cell_shading(cell, shading)
    repeat_table_header(header)

    numeric = _numeric_columns(rows)
    for row_index, values in enumerate(rows):
        cells = table.rows[row_index + 1].cells
        for index in range(len(columns)):
            value = values[index] if index < len(values) else None
            paragraph = cells[index].paragraphs[0]
            if index in numeric:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_run_font(paragraph.add_run(_display(value)), size=_TABLE_SIZE)


def _numeric_columns(rows: list[list[Any]]) -> set[int]:
    """整列都是数的列右对齐：数字左对齐时小数点对不齐，一列数读起来很费劲。"""
    width = max((len(row) for row in rows), default=0)
    numeric: set[int] = set()
    for index in range(width):
        values = [
            row[index] for row in rows if index < len(row) and flatten_text(row[index])
        ]
        if values and all(to_number(value) is not None for value in values):
            numeric.add(index)
    return numeric


def _add_list(document, block: dict[str, Any]) -> bool:
    items = list_items(block)
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run(item))
    return bool(items)


def _add_metrics(document, block: dict[str, Any]) -> bool:
    items = metric_items(block)
    if not items:
        return False
    _add_caption(document, block.get("text"))
    table = create_table(
        document, columns=2, rows=len(items), weights=[30, 70], bordered=False
    )
    for index, item in enumerate(items):
        label_cell, value_cell = table.rows[index].cells
        set_run_font(
            label_cell.paragraphs[0].add_run(flatten_text(item.get("label"))),
            size=_TABLE_SIZE,
            color=_MUTED,
        )
        value = flatten_text(item.get("value"))
        # 中文取值说明它是一段话（产品名、功能单位），不是给眼球抓的数，不加粗
        set_run_font(
            value_cell.paragraphs[0].add_run(value or "—"),
            size=_BODY_SIZE,
            bold=not _CJK.search(value),
        )
    return True


def _bar_cells(row, track_cells: int):
    """一行条形：标签 + 轨道格 + 取值，中间的轨道格交给调用方上色。"""
    cells = row.cells
    return cells[0], cells[1 : 1 + track_cells], cells[1 + track_cells]


def _fill_bar(track: list, *, start: int, length: int, color: str) -> None:
    for offset in range(length):
        position = start + offset
        if 0 <= position < len(track):
            set_cell_shading(track[position], color)


def _bar_row_label(cell, item: dict[str, Any]) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(paragraph.add_run(flatten_text(item.get("label"))), size=_TABLE_SIZE)


def _bar_row_value(cell, item: dict[str, Any], unit: str) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    text = flatten_text(item.get("value"))
    set_run_font(paragraph.add_run(f"{text} {unit}".strip() or "—"), size=_TABLE_SIZE)


def _add_length_bars(
    document, series: list[dict[str, Any]], values: list[float | None], max_abs: float, unit: str
) -> None:
    """同号量值按长度比。整条轨道的底色先铺满，再按比例把前半段染成主色。"""
    track_cells = _TRACK_CELLS
    weights = [_BAR_LABEL_WEIGHT, *([3] * track_cells), _BAR_VALUE_WEIGHT]
    table = create_table(
        document, columns=2 + track_cells, rows=len(series), weights=weights, bordered=False
    )
    for index, item in enumerate(series):
        label_cell, track, value_cell = _bar_cells(table.rows[index], track_cells)
        _bar_row_label(label_cell, item)
        _fill_bar(track, start=0, length=track_cells, color=_TRACK)
        value = values[index]
        if value and max_abs > 0:
            filled = max(1, round(abs(value) / max_abs * track_cells))
            _fill_bar(track, start=0, length=filled, color=_SERIES[0])
        _bar_row_value(value_cell, item, unit)


def _add_diverging_bars(
    document, series: list[dict[str, Any]], values: list[float | None], max_abs: float, unit: str
) -> None:
    """出现负值时零基线居中：正值向右、负值向左，长度仍按绝对值比。

    长度条会把负值画成零宽，读者只看到数字、看不到量级——减排与抵扣在这个领域很常见。
    """
    half = _TRACK_CELLS // 2
    track_cells = half * 2 + 1
    axis = half
    weights = [_BAR_LABEL_WEIGHT, *([3] * track_cells), _BAR_VALUE_WEIGHT]
    table = create_table(
        document, columns=2 + track_cells, rows=len(series), weights=weights, bordered=False
    )
    for index, item in enumerate(series):
        label_cell, track, value_cell = _bar_cells(table.rows[index], track_cells)
        _bar_row_label(label_cell, item)
        _fill_bar(track, start=0, length=track_cells, color=_TRACK)
        set_cell_shading(track[axis], _AXIS)
        value = values[index]
        if value and max_abs > 0:
            length = max(1, round(abs(value) / max_abs * half))
            if value > 0:
                _fill_bar(track, start=axis + 1, length=length, color=_POSITIVE)
            else:
                _fill_bar(track, start=axis - length, length=length, color=_NEGATIVE)
        _bar_row_value(value_cell, item, unit)


def _add_bars(document, block: dict[str, Any]) -> bool:
    series = series_items(block)
    if not series:
        return False
    unit = block_unit(block)
    values = [to_number(item.get("value")) for item in series]
    max_abs = max([abs(value) for value in values if value is not None] or [0])
    _add_caption(document, block.get("text"))
    if any(value is not None and value < 0 for value in values):
        _add_note(document, "正值：排放／增加；负值：减排／抵扣。")
        _add_diverging_bars(document, series, values, max_abs, unit)
    else:
        _add_length_bars(document, series, values, max_abs, unit)
    number_column = f"数值（{unit}）" if unit else "数值"
    _add_data_table(
        document,
        ["分项", number_column],
        [[item.get("label"), item.get("value")] for item in series],
    )
    return True


def _add_share(document, block: dict[str, Any]) -> bool:
    """构成占比：横向堆叠条 + 图例。类目太多、含负值或总量为零时不画占比，直接出表。"""
    series = series_items(block)
    if not series:
        return False
    unit = block_unit(block)
    values = [to_number(item.get("value")) for item in series]
    total = sum(max(value or 0, 0) for value in values)
    number_column = f"数值（{unit}）" if unit else "数值"
    rows = [[item.get("label"), item.get("value")] for item in series]
    _add_caption(document, block.get("text"))

    if any(value is not None and value < 0 for value in values):
        _add_note(document, "含负值分项（减排／抵扣），无法作为占比呈现，改为数据表。")
        _add_data_table(document, ["分项", number_column], rows)
        return True
    if len(series) > _SHARE_SLOTS or total <= 0:
        _add_data_table(document, ["分项", number_column], rows)
        return True

    # 一段一格地染色，最后一段补齐余数，保证总宽度正好铺满
    table = create_table(
        document, columns=_TRACK_CELLS, rows=1, weights=[3] * _TRACK_CELLS, bordered=False
    )
    track = table.rows[0].cells
    cursor = 0
    for index, value in enumerate(values):
        length = (
            round((value or 0) / total * _TRACK_CELLS)
            if index < len(values) - 1
            else _TRACK_CELLS - cursor
        )
        _fill_bar(track, start=cursor, length=length, color=_SERIES[index % _SHARE_SLOTS])
        cursor += length

    # 图例即数据表：色块、分项、数值、占比各一列，读者能把颜色对回分项
    legend = create_table(
        document,
        columns=4,
        rows=len(series) + 1,
        weights=[6, 44, 26, 24],
        bordered=False,
    )
    header = legend.rows[0]
    for index, name in enumerate(("", "分项", number_column, "占比")):
        set_run_font(header.cells[index].paragraphs[0].add_run(name), size=_TABLE_SIZE, bold=True)
    for index, item in enumerate(series):
        cells = legend.rows[index + 1].cells
        set_cell_shading(cells[0], _SERIES[index % _SHARE_SLOTS])
        set_run_font(
            cells[1].paragraphs[0].add_run(flatten_text(item.get("label"))), size=_TABLE_SIZE
        )
        set_run_font(
            cells[2].paragraphs[0].add_run(_display(item.get("value"))), size=_TABLE_SIZE
        )
        set_run_font(
            cells[3].paragraphs[0].add_run(f"{(values[index] or 0) / total * 100:.1f}%"),
            size=_TABLE_SIZE,
        )
    return True


def _add_table_block(document, block: dict[str, Any]) -> bool:
    data = block_data(block)
    rows = table_rows(data)
    if not rows:
        # 没有行数据就退回标题本身，至少不让这块内容凭空消失
        return _add_paragraph(document, block.get("text"))
    columns = table_columns(data)
    width = max(len(row) for row in rows)
    if len(columns) < width:
        # 列名缺失时补空列名，也要出表格：宁可表头留白，也不能退化成带竖线的段落
        columns = [*columns, *([""] * (width - len(columns)))]
    _add_caption(document, block.get("text"))
    _add_data_table(document, columns, rows)
    return True


def _add_raw(document, block: dict[str, Any]) -> None:
    """读不出来内容的块按原样输出：静默丢块会让导出件比在线视图少内容。"""
    payload = _raw_payload(block)
    if payload is None:
        return
    paragraph = _paragraph(document, size=_NOTE_SIZE, indent=False)
    set_run_font(
        paragraph.add_run(_json_text(payload)), size=_NOTE_SIZE, latin=MONO_FONT, color=_MUTED
    )


def _raw_payload(block: dict[str, Any]) -> Any:
    """兜底输出用的原始内容：有 data 用 data，否则用 text；都空则没有可输出的东西。"""
    payload = block.get("data") if block.get("data") is not None else block.get("text")
    if payload is None or payload in ("", {}, []):
        return None
    return payload


def _write_block(document, block: dict[str, Any]) -> bool:
    """按块类型写进文档，返回是否写出了内容。

    判据只有这一处：调用方据此决定兜底，章节标题也据此决定留不留。另写一个
    「这个块算不算空」的谓词，迟早会与真正的渲染分支走岔——读得出内容却排不出来，
    于是内容在纸面上消失而没有任何痕迹。
    """
    block_type = flatten_text(block.get("type"))
    if block_type in ("paragraph", "signature_block", "source_note"):
        written = _add_paragraph(document, block.get("text"))
    elif block_type == "heading":
        written = _add_paragraph(document, block.get("text"), indent=False, style="Heading 2")
    elif block_type == "callout":
        written = _add_callout(document, block)
    elif block_type == "list":
        written = _add_list(document, block)
    elif block_type == "metric_cards":
        written = _add_metrics(document, block)
    elif block_type == "bar_chart":
        written = _add_bars(document, block)
    elif block_type == "donut_chart":
        written = _add_share(document, block)
    elif block_type == "table":
        written = _add_table_block(document, block)
    elif block_type == "page_break":
        document.add_page_break()
        written = True
    else:
        written = False
    if written:
        _add_evidence(document, block.get("evidence_ids"))
    return written


def _render_block(document, block: Any) -> None:
    if not isinstance(block, dict):
        return
    if not _write_block(document, block):
        _add_raw(document, block)


def _render_sections(document, report_ir: dict[str, Any]) -> None:
    for section in report_ir.get("sections") or []:
        if not isinstance(section, dict):
            continue
        blocks = section.get("blocks") or []
        title = flatten_text(section.get("title")) or flatten_text(section.get("section_id"))
        heading = document.add_paragraph(title, style="Heading 1")
        before = len(document.element.body)
        for block in blocks:
            _render_block(document, block)
        # 整节都排不出内容时不留一个孤零零的标题。判据是正文里真的多了几个元素，
        # 而不是另写一套「这个块算不算空」的猜测——那套判断迟早和渲染分支走岔。
        if len(document.element.body) == before:
            heading._element.getparent().remove(heading._element)


def _add_bullets(document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run(item))


def _render_warnings(document, report_ir: dict[str, Any]) -> None:
    warnings = [flatten_text(item) for item in report_ir.get("warnings") or []]
    warnings = [item for item in warnings if item]
    if not warnings:
        return
    document.add_paragraph("阅读提示", style="Heading 1")
    _add_bullets(document, warnings)


def _render_ledger(document, report_ir: dict[str, Any], template: ReportTemplate) -> None:
    ledger = [item for item in report_ir.get("field_ledger") or [] if isinstance(item, dict)]
    if not ledger:
        return
    labels = {
        field.get("id"): field.get("label") or field.get("id")
        for field in template.definition.get("fields") or []
    }

    def name_of(item: dict[str, Any]) -> str:
        field_id = flatten_text(item.get("field_id"))
        return flatten_text(labels.get(field_id)) or field_id

    def status_of(item: dict[str, Any]) -> str:
        status = flatten_text(item.get("status"))
        return FIELD_STATUS_LABELS.get(status, status)

    document.add_paragraph("字段台账", style="Heading 1")
    pending = [item for item in ledger if item.get("status") in UNRESOLVED_STATUSES]
    if pending:
        _add_paragraph(
            document,
            f"以下 {len(pending)} 项在现有材料中未能落实，正文相应位置按资料缺口处理。",
            indent=False,
        )
        for item in pending:
            notes = "；".join(flatten_text(note) for note in item.get("notes") or [])
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(0)
            set_run_font(paragraph.add_run(name_of(item)), bold=True)
            set_run_font(paragraph.add_run(f"　{status_of(item)}"), size=_NOTE_SIZE, color=_MUTED)
            if notes:
                set_run_font(paragraph.add_run(f"　{notes}"), size=_NOTE_SIZE, color=_MUTED)

    rows = [
        [
            name_of(item),
            status_of(item),
            item.get("value"),
            item.get("unit"),
            "、".join(flatten_text(eid) for eid in item.get("evidence_ids") or []),
        ]
        for item in ledger
    ]
    _add_data_table(document, ["字段", "状态", "取值", "单位", "证据"], rows)


def _render_evidence(document, report_ir: dict[str, Any]) -> None:
    evidence = [item for item in report_ir.get("evidence") or [] if isinstance(item, dict)]
    if not evidence:
        return
    document.add_paragraph("证据台账", style="Heading 1")
    rows = [
        [
            item.get("evidence_id"),
            item.get("source_type"),
            evidence_location(item),
            trim_text(item.get("excerpt"), 200),
        ]
        for item in evidence
    ]
    _add_data_table(document, ["证据", "来源", "位置", "摘录"], rows)


def _render_calculations(document, report_ir: dict[str, Any]) -> None:
    calculations = [
        item for item in report_ir.get("calculations") or [] if isinstance(item, dict)
    ]
    if not calculations:
        return
    document.add_paragraph("计算过程", style="Heading 1")
    for item in calculations:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(0)
        # 公式编号用等宽字体标出来，但不加反引号：那是 Markdown 的写法，Word 里只会多两个字符
        set_run_font(
            paragraph.add_run(flatten_text(item.get("formula_id"))), latin=MONO_FONT, bold=True
        )
        set_run_font(
            paragraph.add_run(
                f" v{flatten_text(item.get('formula_version'))}："
                f"{flatten_text(item.get('output_field_id'))} = "
                f"{flatten_text(item.get('value'))} {flatten_text(item.get('unit'))}".rstrip()
            )
        )
        _add_note(document, f"输入：{_json_text(item.get('inputs'))}")


def _render_limitations(document, report_ir: dict[str, Any]) -> None:
    limitations = [flatten_text(item) for item in report_ir.get("limitations") or []]
    limitations = [item for item in limitations if item]
    if not limitations:
        return
    document.add_paragraph("使用限制", style="Heading 1")
    _add_bullets(document, limitations)


def _add_footer(document, run: ReportRun) -> None:
    paragraph = document.sections[0].footer.paragraphs[0]
    add_page_number(
        paragraph,
        prefix=f"本文件由 EnerLedger 生成　报告任务 {str(run.id)[:8]}　第 ",
        suffix=" 页",
    )


def _render_cover(document, *, run: ReportRun, name: str) -> None:
    """封面只放读者认得的东西：报告名称与生成日期。

    任务 UUID、模板 slug、``文档版本 vNone`` 这类内部标识不进正文——直传材料没有
    文档版本，硬写出来就是一行「vNone」摆在封面上。
    """
    eyebrow = _paragraph(document, size=_TABLE_SIZE, indent=False)
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(eyebrow.add_run("报告"), size=_TABLE_SIZE, color=_ACCENT, east_asia=COVER_FONT)

    title = _paragraph(document, size=_COVER_TITLE_SIZE, indent=False)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        title.add_run(name), size=_COVER_TITLE_SIZE, bold=True, east_asia=COVER_FONT
    )

    stamp = run.finished_at or run.created_at
    if stamp is not None:
        subtitle = _paragraph(document, size=_NOTE_SIZE, indent=False)
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(
            subtitle.add_run(f"生成日期 {stamp:%Y-%m-%d}"), size=_NOTE_SIZE, color=_MUTED
        )
    _paragraph(document, size=_NOTE_SIZE, indent=False)


def render_report_docx(
    *, run: ReportRun, template: ReportTemplate, report_ir: dict[str, Any]
) -> bytes:
    """把 ReportIR 渲染成 Word 字节流。"""
    if not isinstance(report_ir, dict):
        report_ir = {}
    name = flatten_text(template.definition.get("name")) or flatten_text(run.template_id)

    document = WordDocument()
    configure_document_styles(document)
    set_page_setup(document)
    enable_field_update(document)
    document.core_properties.title = name
    document.core_properties.subject = flatten_text(run.report_type)
    document.core_properties.author = "EnerLedger"

    _render_cover(document, run=run, name=name)
    # 阅读提示放在正文之前、使用限制放在正文之后，与报告详情页的阅读顺序一致
    _render_warnings(document, report_ir)
    _render_sections(document, report_ir)
    _render_ledger(document, report_ir, template)
    _render_evidence(document, report_ir)
    _render_calculations(document, report_ir)
    _render_limitations(document, report_ir)
    _add_footer(document, run)

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()

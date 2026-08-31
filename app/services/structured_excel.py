"""Known Excel template profiling and deterministic EPA factor extraction."""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

EPA_SHEET = "Emission Factors Hub"
EPA_TEMPLATE_CODE = "EPA_GHG_EMISSION_FACTORS_HUB"
UNCERTAINTY_TEMPLATE_CODE = "GHG_PRODUCT_UNCERTAINTY_TEMPLATE"

EPA_TABLES = {
    1: ("stationary_combustion", "Stationary Combustion"),
    2: ("mobile_combustion_co2", "Mobile Combustion CO2"),
    3: ("mobile_onroad_gasoline", "On-Road Gasoline Vehicles"),
    4: ("mobile_onroad_alternative", "On-Road Diesel and Alternative Fuel Vehicles"),
    5: ("mobile_nonroad", "Non-Road Vehicles"),
    6: ("electricity", "Electricity"),
    7: ("steam_heat", "Steam and Heat"),
    8: ("scope3_transportation", "Scope 3 Transportation and Distribution"),
    9: ("scope3_waste", "Scope 3 Waste and End-of-Life"),
    10: ("scope3_travel_commuting", "Scope 3 Business Travel and Employee Commuting"),
    11: ("gwp", "Global Warming Potential"),
    12: ("blended_refrigerant_gwp", "Blended Refrigerant GWP"),
}

PARQUET_SCHEMA: dict[str, str] = {
    "edition_year": "INTEGER",
    "table_code": "VARCHAR",
    "scope": "VARCHAR",
    "category": "VARCHAR",
    "activity_name": "VARCHAR",
    "fuel_type": "VARCHAR",
    "vehicle_type": "VARCHAR",
    "material_type": "VARCHAR",
    "treatment_method": "VARCHAR",
    "region_code": "VARCHAR",
    "region_name": "VARCHAR",
    "model_year_raw": "VARCHAR",
    "gas": "VARCHAR",
    "factor_value": "DECIMAL(38, 12)",
    "numerator_unit": "VARCHAR",
    "denominator_unit": "VARCHAR",
    "factor_basis": "VARCHAR",
    "lifecycle_boundary": "VARCHAR",
    "applicability": "VARCHAR",
    "source_sheet": "VARCHAR",
    "source_row": "INTEGER",
    "source_file": "VARCHAR",
    "source_note": "VARCHAR",
}


class StructuredWorkbookError(ValueError):
    """Workbook is unsupported or violates a known template contract."""


@dataclass(frozen=True)
class WorkbookProfile:
    template_code: str
    asset_code: str
    asset_name: str
    asset_type: str
    version_label: str
    edition_year: int | None
    sheet_names: list[str]
    table_count: int
    formula_count: int
    details: dict[str, Any]


@dataclass(frozen=True)
class ParsedTable:
    table_code: str
    display_name: str
    source_sheet: str
    source_range: str
    records: list[dict[str, Any]]


@dataclass(frozen=True)
class ParsedWorkbook:
    profile: WorkbookProfile
    tables: list[ParsedTable]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _decimal(value: Any) -> Decimal | None:
    text = _text(value)
    if not text or text.upper() in {"NA", "N/A", "-"}:
        return None
    try:
        normalized = re.sub(r"^[<>~≈≥≤]+\s*", "", text).replace(",", "")
        return Decimal(normalized)
    except InvalidOperation as exc:
        raise StructuredWorkbookError(f"无法把 {text!r} 转换为数值") from exc


def _maybe_decimal(value: Any) -> Decimal | None:
    try:
        return _decimal(value)
    except StructuredWorkbookError:
        return None


def _formula_count(workbook) -> int:
    count = 0
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            count += sum(
                1 for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")
            )
    return count


def _edition_year(filename: str) -> int | None:
    match = re.search(r"(?<!\d)(20\d{2})(?!\d)", filename)
    return int(match.group(1)) if match else None


def profile_workbook(path: Path, filename: str) -> WorkbookProfile:
    """Classify a workbook without trusting filename alone."""

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet_names = list(workbook.sheetnames)
        formulas = _formula_count(workbook)
        if EPA_SHEET in sheet_names:
            sheet = workbook[EPA_SHEET]
            found: dict[int, int] = {}
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    match = re.fullmatch(r"Table\s+(\d+)", _text(value), flags=re.IGNORECASE)
                    if match:
                        found[int(match.group(1))] = len(found) + 1
            if set(found) != set(EPA_TABLES):
                missing = sorted(set(EPA_TABLES) - set(found))
                raise StructuredWorkbookError(f"EPA 工作簿缺少逻辑表: {missing}")
            year = _edition_year(filename)
            if year is None:
                raise StructuredWorkbookError("EPA 工作簿文件名中缺少四位版本年份")
            return WorkbookProfile(
                template_code=EPA_TEMPLATE_CODE,
                asset_code="epa_ghg_emission_factors",
                asset_name="EPA GHG Emission Factors Hub",
                asset_type="REFERENCE_DATA",
                version_label=str(year),
                edition_year=year,
                sheet_names=sheet_names,
                table_count=12,
                formula_count=formulas,
                details={"logical_tables": [EPA_TABLES[index][0] for index in sorted(EPA_TABLES)]},
            )
        required = {"Instruction", "Notes", "Info", "Calculations"}
        if required.issubset(sheet_names):
            calculations = workbook["Calculations"]
            if calculations.max_column < 82:
                raise StructuredWorkbookError("不确定性计算工作表缺少预期计算列")
            return WorkbookProfile(
                template_code=UNCERTAINTY_TEMPLATE_CODE,
                asset_code="ghg_product_uncertainty",
                asset_name="产品温室气体清单不确定性模型",
                asset_type="CALCULATION_MODEL",
                version_label="v1",
                edition_year=None,
                sheet_names=sheet_names,
                table_count=4,
                formula_count=formulas,
                details={
                    "calculation_rows": calculations.max_row,
                    "calculation_columns": calculations.max_column,
                    "activation_state": "MODEL_VALIDATION_REQUIRED",
                },
            )
        raise StructuredWorkbookError("暂不支持该 Excel 模板")
    finally:
        workbook.close()


def _find_table_rows(rows: list[tuple[Any, ...]]) -> dict[int, int]:
    result: dict[int, int] = {}
    for index, row in enumerate(rows, start=1):
        first = _text(row[0] if row else None)
        match = re.fullmatch(r"Table\s+(\d+)", first, flags=re.IGNORECASE)
        if match:
            result[int(match.group(1))] = index
    if set(result) != set(EPA_TABLES):
        raise StructuredWorkbookError("EPA 逻辑表数量或编号不符合预期")
    return result


def _trim_leading_empty_columns(rows: list[tuple[Any, ...]]) -> list[tuple[Any, ...]]:
    populated = [index for row in rows for index, value in enumerate(row) if _text(value)]
    if not populated:
        return rows
    offset = min(populated)
    return [tuple(row[offset:]) for row in rows]


def _find_header(
    rows: list[tuple[Any, ...]], start: int, end: int, needle: str | tuple[str, ...]
) -> int:
    needles = (needle,) if isinstance(needle, str) else needle
    for row_number in range(start + 1, min(start + 8, end)):
        cells = [_text(value).lower() for value in rows[row_number - 1]]
        if any(cell.startswith(candidate.lower()) for candidate in needles for cell in cells):
            return row_number
    raise StructuredWorkbookError(f"在 Table 起始行 {start} 后未找到表头 {needle!r}")


def _table_note(rows: list[tuple[Any, ...]], data_start: int, end: int) -> str:
    notes: list[str] = []
    started = False
    for row_number in range(data_start, end):
        values = [_text(value) for value in rows[row_number - 1]]
        meaningful = [value for value in values if value]
        if not meaningful:
            continue
        lead = meaningful[0]
        if lead.startswith(("Source:", "Notes:", "Note:", "http://", "https://")):
            started = True
        if started:
            notes.append(" ".join(meaningful))
    return "\n".join(notes)[:16000]


def _base_record(
    *,
    year: int,
    table_code: str,
    source_row: int,
    source_file: str,
    source_note: str,
    **values: Any,
) -> dict[str, Any]:
    record = {key: None for key in PARQUET_SCHEMA}
    record.update(
        {
            "edition_year": year,
            "table_code": table_code,
            "applicability": "AVAILABLE",
            "source_sheet": EPA_SHEET,
            "source_row": source_row,
            "source_file": source_file,
            "source_note": source_note or None,
        }
    )
    record.update(values)
    return record


def _append_gases(
    records: list[dict[str, Any]],
    *,
    gases: list[tuple[str, Any]],
    numerator_units: dict[str, str],
    denominator_unit: str,
    common: dict[str, Any],
) -> None:
    for gas, raw_value in gases:
        value = _decimal(raw_value)
        if value is None:
            continue
        records.append(
            _base_record(
                **common,
                gas=gas,
                factor_value=value,
                numerator_unit=numerator_units[gas],
                denominator_unit=denominator_unit,
            )
        )


def _native_fuel_unit(category: str) -> str:
    lowered = category.lower()
    if "gaseous" in lowered or "gas" in lowered:
        return "scf"
    if "liquid" in lowered or "petroleum" in lowered:
        return "gallon"
    return "short ton"


def _parse_table(
    number: int,
    rows: list[tuple[Any, ...]],
    start: int,
    end: int,
    *,
    year: int,
    filename: str,
) -> ParsedTable:
    table_code, display_name = EPA_TABLES[number]
    header_needles = {
        1: "Fuel Type",
        2: "Fuel Type",
        3: "Vehicle Type",
        4: "Vehicle Type",
        5: "Vehicle Type",
        6: "eGRID Subregion",
        7: "CO2 Factor",
        8: "Vehicle Type",
        9: "Material",
        10: "Vehicle Type",
        11: ("Industrial Designation", "Gas"),
        12: "ASHRAE",
    }
    header = _find_header(rows, start, end, header_needles[number])
    data_start = header + (2 if number in {1, 6} else 1)
    note = _table_note(rows, data_start, end)
    records: list[dict[str, Any]] = []
    category = ""
    vehicle = ""
    fuel = ""

    for row_number in range(data_start, end):
        row = list(rows[row_number - 1]) + [None] * 12
        b, c, d, e, f, g, h, i, j = row[1:10]
        lead = _text(b)
        if lead.startswith(("Source:", "Notes:", "Note:")):
            break
        common = {
            "year": year,
            "table_code": table_code,
            "source_row": row_number,
            "source_file": filename,
            "source_note": note,
        }
        if number == 1:
            numeric = [_maybe_decimal(value) for value in (c, d, e, f, g, h, i)]
            if lead and all(value is None for value in numeric):
                category = lead
                continue
            if not lead or all(value is None for value in numeric):
                continue
            native_unit = _native_fuel_unit(category)
            if numeric[0] is not None:
                records.append(
                    _base_record(
                        **common,
                        category=category or None,
                        activity_name=lead,
                        fuel_type=lead,
                        gas="HEAT_CONTENT",
                        factor_value=numeric[0],
                        numerator_unit="mmBtu",
                        denominator_unit=native_unit,
                        factor_basis="HHV",
                        lifecycle_boundary="COMBUSTION_ONLY",
                    )
                )
            _append_gases(
                records,
                gases=list(zip(("CO2", "CH4", "N2O"), (d, e, f), strict=True)),
                numerator_units={"CO2": "kg CO2", "CH4": "g CH4", "N2O": "g N2O"},
                denominator_unit="mmBtu",
                common={
                    **common,
                    "category": category or None,
                    "activity_name": lead,
                    "fuel_type": lead,
                    "factor_basis": "HHV",
                    "lifecycle_boundary": "COMBUSTION_ONLY",
                },
            )
            _append_gases(
                records,
                gases=list(zip(("CO2", "CH4", "N2O"), (g, h, i), strict=True)),
                numerator_units={"CO2": "kg CO2", "CH4": "g CH4", "N2O": "g N2O"},
                denominator_unit=native_unit,
                common={
                    **common,
                    "category": category or None,
                    "activity_name": lead,
                    "fuel_type": lead,
                    "factor_basis": "HHV",
                    "lifecycle_boundary": "COMBUSTION_ONLY",
                },
            )
        elif number == 2:
            value = _decimal(c)
            if lead and value is not None:
                records.append(
                    _base_record(
                        **common,
                        activity_name=lead,
                        fuel_type=lead,
                        gas="CO2",
                        factor_value=value,
                        numerator_unit="kg CO2",
                        denominator_unit=_text(d),
                        lifecycle_boundary="TANK_TO_WHEEL",
                    )
                )
        elif number == 3:
            vehicle = lead or vehicle
            if not vehicle or not _text(c):
                continue
            _append_gases(
                records,
                gases=[("CH4", d), ("N2O", e)],
                numerator_units={"CH4": "g CH4", "N2O": "g N2O"},
                denominator_unit="vehicle-mile",
                common={
                    **common,
                    "activity_name": vehicle,
                    "vehicle_type": vehicle,
                    "model_year_raw": _text(c),
                    "lifecycle_boundary": "TANK_TO_WHEEL",
                },
            )
        elif number == 4:
            vehicle = lead or vehicle
            fuel = _text(c) or fuel
            if not vehicle or not fuel or not _text(d):
                continue
            _append_gases(
                records,
                gases=[("CH4", e), ("N2O", f)],
                numerator_units={"CH4": "g CH4", "N2O": "g N2O"},
                denominator_unit="vehicle-mile",
                common={
                    **common,
                    "activity_name": vehicle,
                    "vehicle_type": vehicle,
                    "fuel_type": fuel,
                    "model_year_raw": _text(d),
                    "lifecycle_boundary": "TANK_TO_WHEEL",
                },
            )
        elif number == 5:
            vehicle = lead or vehicle
            fuel = _text(c)
            if not vehicle or not fuel:
                continue
            _append_gases(
                records,
                gases=[("CH4", d), ("N2O", e)],
                numerator_units={"CH4": "g CH4", "N2O": "g N2O"},
                denominator_unit="gallon",
                common={
                    **common,
                    "activity_name": vehicle,
                    "vehicle_type": vehicle,
                    "fuel_type": fuel,
                    "lifecycle_boundary": "TANK_TO_WHEEL",
                },
            )
        elif number == 6:
            region_code, region_name = lead, _text(c)
            if not region_code or not region_name:
                continue
            for basis, values in (("TOTAL_OUTPUT", (d, e, f)), ("NON_BASELOAD", (g, h, i))):
                _append_gases(
                    records,
                    gases=list(zip(("CO2", "CH4", "N2O"), values, strict=True)),
                    numerator_units={"CO2": "lb CO2", "CH4": "lb CH4", "N2O": "lb N2O"},
                    denominator_unit="MWh",
                    common={
                        **common,
                        "activity_name": region_name,
                        "region_code": region_code,
                        "region_name": region_name,
                        "factor_basis": basis,
                        "lifecycle_boundary": "GENERATION_COMBUSTION_ONLY",
                    },
                )
            loss = _decimal(j)
            if loss is not None:
                records.append(
                    _base_record(
                        **common,
                        activity_name=region_name,
                        region_code=region_code,
                        region_name=region_name,
                        gas="GRID_LOSS",
                        factor_value=loss,
                        numerator_unit="ratio",
                        denominator_unit="electricity",
                        factor_basis="T_AND_D_LOSS",
                    )
                )
        elif number == 7:
            if not lead:
                continue
            _append_gases(
                records,
                gases=[("CO2", c), ("CH4", d), ("N2O", e)],
                numerator_units={"CO2": "kg CO2", "CH4": "g CH4", "N2O": "g N2O"},
                denominator_unit="mmBtu",
                common={
                    **common,
                    "activity_name": lead,
                    "factor_basis": "PURCHASED_STEAM_HEAT",
                    "lifecycle_boundary": "COMBUSTION_ONLY",
                },
            )
        elif number in {8, 10}:
            if not lead:
                continue
            _append_gases(
                records,
                gases=[("CO2", c), ("CH4", d), ("N2O", e)],
                numerator_units={"CO2": "kg CO2", "CH4": "g CH4", "N2O": "g N2O"},
                denominator_unit=_text(f),
                common={
                    **common,
                    "scope": "SCOPE_3",
                    "activity_name": lead,
                    "vehicle_type": lead,
                    "factor_basis": "DISTANCE_BASED",
                    "lifecycle_boundary": "TANK_TO_WHEEL",
                },
            )
        elif number == 9:
            if not lead:
                continue
            headers = rows[header - 1]
            for column_index, raw in enumerate((c, d, e, f, g, h), start=2):
                method = _text(headers[column_index])
                raw_text = _text(raw)
                if not method or not raw_text:
                    continue
                value = _decimal(raw)
                records.append(
                    _base_record(
                        **common,
                        scope="SCOPE_3",
                        activity_name=lead,
                        material_type=lead,
                        treatment_method=method,
                        gas="CO2E",
                        factor_value=value,
                        numerator_unit="metric ton CO2e",
                        denominator_unit="short ton material",
                        factor_basis="WASTE_TYPE_SPECIFIC",
                        applicability="AVAILABLE" if value is not None else "NOT_APPLICABLE",
                    )
                )
        elif number == 11:
            formula = _text(c) if _decimal(d) is not None else lead
            value = _decimal(d) if _decimal(d) is not None else _decimal(c)
            if lead and formula and value is not None:
                records.append(
                    _base_record(
                        **common,
                        activity_name=lead,
                        gas=formula,
                        factor_value=value,
                        numerator_unit="kg CO2e",
                        denominator_unit=f"kg {formula}",
                        factor_basis=(
                            "100_YEAR_GWP_LOWER_BOUND"
                            if _text(c).startswith(">")
                            else "100_YEAR_GWP"
                        ),
                    )
                )
        elif number == 12:
            if lead and _decimal(c) is not None:
                records.append(
                    _base_record(
                        **common,
                        activity_name=lead,
                        gas=lead,
                        factor_value=_decimal(c),
                        numerator_unit="kg CO2e",
                        denominator_unit=f"kg {lead}",
                        factor_basis="100_YEAR_GWP",
                        category=_text(d) or None,
                    )
                )
    if not records:
        raise StructuredWorkbookError(f"{display_name} 未提取到有效记录")
    last_row = max(record["source_row"] for record in records)
    return ParsedTable(
        table_code=table_code,
        display_name=display_name,
        source_sheet=EPA_SHEET,
        source_range=f"A{start}:J{last_row}",
        records=records,
    )


def parse_epa_workbook(path: Path, filename: str) -> ParsedWorkbook:
    profile = profile_workbook(path, filename)
    if profile.template_code != EPA_TEMPLATE_CODE or profile.edition_year is None:
        raise StructuredWorkbookError("该文件不是 EPA 排放因子模板")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = _trim_leading_empty_columns(list(workbook[EPA_SHEET].iter_rows(values_only=True)))
    finally:
        workbook.close()
    starts = _find_table_rows(rows)
    tables: list[ParsedTable] = []
    for number in sorted(EPA_TABLES):
        start = starts[number]
        end = starts.get(number + 1, len(rows) + 1)
        tables.append(
            _parse_table(
                number,
                rows,
                start,
                end,
                year=profile.edition_year,
                filename=filename,
            )
        )
    return ParsedWorkbook(profile=profile, tables=tables)


def write_parquet_tables(parsed: ParsedWorkbook, output_dir: Path) -> dict[str, Path]:
    """Write one immutable Parquet object per logical table using DuckDB."""

    import duckdb

    output_dir.mkdir(parents=True, exist_ok=True)
    columns = list(PARQUET_SCHEMA)
    definitions = ", ".join(f'"{name}" {type_name}' for name, type_name in PARQUET_SCHEMA.items())
    placeholders = ", ".join("?" for _ in columns)
    outputs: dict[str, Path] = {}
    for table in parsed.tables:
        destination = output_dir / f"{table.table_code}.parquet"
        connection = duckdb.connect(":memory:")
        try:
            connection.execute(f"CREATE TABLE factors ({definitions})")
            connection.executemany(
                f"INSERT INTO factors VALUES ({placeholders})",
                [[record.get(column) for column in columns] for record in table.records],
            )
            escaped = str(destination).replace("'", "''")
            connection.execute(f"COPY factors TO '{escaped}' (FORMAT PARQUET, COMPRESSION ZSTD)")
        finally:
            connection.close()
        outputs[table.table_code] = destination
    return outputs
